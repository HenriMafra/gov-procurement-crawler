"""
Analisa o Excel gerado e reporta: abas, colunas, exemplos de dados,
campos vazios, valores zerados, status N/D, etc.
"""
import sys
sys.path.insert(0, ".")

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook

PATH = r"C:\temp_b2g\outputs\coleta_orgaos_especificos_2026-06-30.xlsx"

wb = load_workbook(PATH, read_only=True, data_only=True)
print(f"Abas encontradas: {wb.sheetnames}\n")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"[{sheet_name}] Vazia.\n")
        continue

    headers = rows[0]
    data_rows = rows[1:]
    print(f"{'='*60}")
    print(f"ABA: {sheet_name} | Total linhas de dados: {len(data_rows)}")
    print(f"Colunas ({len(headers)}): {list(headers)}")

    if data_rows:
        # Analisa campos vazios e valores zerados
        zero_valor = 0
        nd_status = 0
        sem_fornecedor = 0
        sem_objeto = 0

        # Tenta encontrar índices das colunas relevantes
        h_lower = [str(h).lower() if h else "" for h in headers]
        idx_valor = next((i for i, h in enumerate(h_lower) if "valor" in h), None)
        idx_status = next((i for i, h in enumerate(h_lower) if "status" in h), None)
        idx_forn = next((i for i, h in enumerate(h_lower) if "fornecedor" in h and "cnpj" not in h), None)
        idx_obj = next((i for i, h in enumerate(h_lower) if "objeto" in h), None)

        for row in data_rows:
            if idx_valor is not None:
                v = row[idx_valor]
                if v is None or v == 0 or v == 0.0:
                    zero_valor += 1
            if idx_status is not None:
                s = row[idx_status]
                if s in (None, "N/D", ""):
                    nd_status += 1
            if idx_forn is not None:
                f = row[idx_forn]
                if not f:
                    sem_fornecedor += 1
            if idx_obj is not None:
                o = row[idx_obj]
                if not o:
                    sem_objeto += 1

        if idx_valor is not None:
            print(f"  >> Valores R$=0 ou nulos:  {zero_valor}/{len(data_rows)} ({100*zero_valor//max(len(data_rows),1)}%)")
        if idx_status is not None:
            print(f"  >> Status N/D ou nulos:    {nd_status}/{len(data_rows)} ({100*nd_status//max(len(data_rows),1)}%)")
        if idx_forn is not None:
            print(f"  >> Fornecedor vazio:       {sem_fornecedor}/{len(data_rows)}")
        if idx_obj is not None:
            print(f"  >> Objeto vazio:           {sem_objeto}/{len(data_rows)}")

        # Mostra 3 exemplos
        print(f"\n  Exemplos (3 primeiras linhas de dados):")
        for i, row in enumerate(data_rows[:3]):
            print(f"  Linha {i+2}: {dict(zip(headers, row))}")
    print()

wb.close()
print("Análise concluída.")
