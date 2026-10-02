# -*- coding: utf-8 -*-
"""
Coleta TODOS os contratos da ADASA via Portal da Transparência do GDF.
API descoberta via engenharia reversa da SPA Angular (SIGGO/SEEC):
  https://www.transparencia.df.gov.br/api/licitacoes-contratos/contrato
  Parâmetro chave: listaCodigoUnidadeGestora=150206 (código da ADASA)
Substitui os 32 contratos parciais vindos do cache PNCP.
"""
import urllib.request, ssl, json, re, datetime, time
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ctx = ssl.create_default_context()
ctx.check_hostname = False
ctx.verify_mode = ssl.CERT_NONE
UA = {'User-Agent': 'Mozilla/5.0', 'Accept': 'application/json'}
BASE = 'https://www.transparencia.df.gov.br/api/licitacoes-contratos/contrato'
EXCEL_PATH = r"C:\Users\henri\Desktop\coleta_orgaos_especificos_2026-06-30.xlsx"
ADASA_UG = '150206'
HOJE = datetime.date.today()

ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
def san(v):
    if isinstance(v, str):
        return ILLEGAL.sub("", v).strip()
    return v

def get_json(url, timeout=20):
    try:
        req = urllib.request.Request(url, headers=UA)
        r = urllib.request.urlopen(req, timeout=timeout, context=ctx)
        return json.loads(r.read().decode('utf-8', errors='replace'))
    except Exception as e:
        print(f"  ERRO {url[-80:]}: {e}")
        return {}

def normalizar_data(s):
    if not s:
        return ""
    s = str(s).strip()
    try:
        return datetime.datetime.strptime(s, "%d/%m/%Y").strftime("%Y-%m-%d")
    except:
        return s

def processar_item(item):
    dt_ini = normalizar_data(item.get("dataInicio"))
    dt_fim = normalizar_data(item.get("dataFim"))
    valor = float(item.get("valorContrato") or 0)

    status = "N/D"
    if dt_fim:
        try:
            df_date = datetime.datetime.strptime(dt_fim[:10], "%Y-%m-%d").date()
            status = "Vigente" if df_date >= HOJE else "Expirado"
        except:
            pass

    return {
        "grupo": "ADASA", "sigla": "ADASA",
        "orgao": "ADASA - Agência Reguladora de Águas e Saneamento do DF",
        "cnpj_orgao": "02.413.606/0001-09",
        "unidade": san(item.get("unidadeGestora") or ""),
        "num_contrato": san(item.get("numeroOriginal") or item.get("numeroContrato") or ""),
        "fornecedor": san(item.get("credor") or ""),
        "cnpj_fornecedor": san(item.get("codigoCredorMascarado") or ""),
        "objeto": san(item.get("objeto") or ""),
        "valor": valor,
        "dt_assinatura": "", "dt_inicio": dt_ini, "dt_fim": dt_fim,
        "status": status,
        "modalidade": san(item.get("especie") or ""),
        "link": f"https://www.transparencia.df.gov.br/#/licitacoes-contratos/contratos",
        "fonte": "Portal Transparência GDF (SIGGO)",
    }

def coletar_todos():
    contratos = {}
    anos = list(range(1998, 2027))
    for ano in anos:
        pagina = 0
        while True:
            url = f"{BASE}?anoInicio={ano}&listaCodigoUnidadeGestora={ADASA_UG}&page={pagina}&size=100"
            r = get_json(url)
            content = r.get("content", [])
            if not content:
                break
            for item in content:
                chave = item.get("numeroContrato", "") + "_" + item.get("numeroOriginal", "")
                if chave not in contratos:
                    contratos[chave] = processar_item(item)
            total_pages = r.get("totalPages", 1)
            pagina += 1
            if pagina >= total_pages:
                break
            time.sleep(0.2)
        if content if 'content' in dir() else False:
            pass
        time.sleep(0.3)
    return list(contratos.values())

def adicionar_aba_adasa(contratos):
    wb = load_workbook(EXCEL_PATH)

    fh = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fb = Font(name="Segoe UI", size=10)
    al = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ac = Alignment(horizontal="center", vertical="center")
    ar = Alignment(horizontal="right", vertical="center")
    fill_hdr = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_vg  = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_ex  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    brd = Border(left=Side(style="thin",color="D0D0D0"),right=Side(style="thin",color="D0D0D0"),
                 top=Side(style="thin",color="D0D0D0"),bottom=Side(style="thin",color="D0D0D0"))
    HDRS = ["Grupo","Sigla","Orgao Contratante","CNPJ Orgao","Unidade",
            "Num Contrato","Fornecedor Contratado","CNPJ Fornecedor",
            "Objeto do Contrato","Valor Total (R$)","Assinatura",
            "Inicio Vigencia","Fim Vigencia","Status","Modalidade","Link"]
    WIDTHS = [10,10,42,18,30,15,35,18,60,16,13,13,13,10,20,50]

    def write_header(ws):
        for ci, h in enumerate(HDRS, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = fh; c.alignment = ac; c.fill = fill_hdr
        ws.row_dimensions[1].height = 22

    def write_row(ws, ri, r):
        vals = [r["grupo"],r["sigla"],r["orgao"],r["cnpj_orgao"],r["unidade"],
                r["num_contrato"],r["fornecedor"],r["cnpj_fornecedor"],
                r["objeto"],r["valor"],r["dt_assinatura"],r["dt_inicio"],
                r["dt_fim"],r["status"],r["modalidade"],r["link"]]
        fr = fill_vg if r["status"]=="Vigente" else (fill_ex if r["status"]=="Expirado" else None)
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = fb; c.border = brd
            c.alignment = ar if ci==10 else (al if ci in (3,5,7,9,15,16) else ac)
            if ci == 10: c.number_format = "R$ #,##0.00"
            if fr: c.fill = fr
        ws.row_dimensions[ri].height = 30

    if "ADASA" in wb.sheetnames:
        del wb["ADASA"]
    ws_a = wb.create_sheet("ADASA")
    write_header(ws_a)
    rows_sorted = sorted(contratos, key=lambda x: -x["valor"])
    for ri, r in enumerate(rows_sorted, 2):
        write_row(ws_a, ri, r)
    for ci, w in enumerate(WIDTHS, 1):
        ws_a.column_dimensions[get_column_letter(ci)].width = w

    # Remover linhas antigas de ADASA em "Todos os Contratos" e reinserir
    ws_all = wb["Todos os Contratos"]
    rows_to_delete = []
    for ri in range(2, ws_all.max_row + 1):
        if ws_all.cell(row=ri, column=1).value == "ADASA":
            rows_to_delete.append(ri)
    for ri in reversed(rows_to_delete):
        ws_all.delete_rows(ri)

    nxt = ws_all.max_row + 1
    for r in rows_sorted:
        write_row(ws_all, nxt, r)
        nxt += 1

    wb.save(EXCEL_PATH)
    print(f"Excel salvo: {EXCEL_PATH}")

if __name__ == "__main__":
    print("=== Coleta ADASA via Portal Transparência GDF ===")
    contratos = coletar_todos()
    vg = sum(1 for r in contratos if r["status"] == "Vigente")
    val = sum(r["valor"] for r in contratos)
    print(f"\nTotal ADASA: {len(contratos)} contratos ({vg} vigentes), R$ {val:,.0f}")
    if contratos:
        adicionar_aba_adasa(contratos)
    print("=== CONCLUIDO ===")
