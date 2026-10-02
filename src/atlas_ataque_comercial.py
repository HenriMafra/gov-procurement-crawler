# -*- coding: utf-8 -*-
"""
ATLAS B2G — Lista de Ataque Comercial Imediato (modo "ataque")
==============================================================
Reaproveita o pipeline validado (atlas_pncp_ingest.py): coleta PNCP + cache +
classificador de TI (com tratamento de 'servidor público') + normalização.
Adiciona foco em JANELA COMERCIAL (vencidos recentes / vencendo ≤180d), score
comercial refinado, argumento comercial, urgência e próxima ação específica.

Gera:
  atlas_lista_ataque_comercial.csv        (base completa, ordenada por score)
  atlas_lista_ataque_comercial.xlsx       (executivo: abas + badges + filtros)
  atlas_top_oportunidades_para_vendedores.csv (1 linha/oportunidade, linguagem prática)
  atlas_relatorio_ataque_comercial.md     (relatório executivo da rodada)
  atlas_pncp_log_execucao_ataque.txt      (log + parâmetros + tempo)

Rodar:  python atlas_ataque_comercial.py     (W1 ~24m vem do cache; W2 ~13m é coletada)
"""
import sys, os, csv, time, datetime, collections
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
import atlas_pncp_ingest as core   # reuso do pipeline

HOJE = core.HOJE
def ymd(d): return d.strftime("%Y%m%d")
def dminus(n): return ymd(HOJE - datetime.timedelta(days=n))

# =============================== CONFIG (modo ataque) ===============================
ATK = {
    "UFS": ["DF", "GO"],
    # Janelas de PUBLICAÇÃO escolhidas para capturar contratos que vencem AGORA:
    #  - ~24 meses atrás: contratos de 24m encerrando agora (cacheado da rodada anterior)
    #  - ~13 meses atrás: contratos de 12m encerrando agora / vencidos recentemente
    "JANELAS": [(dminus(744), dminus(730)), (dminus(411), dminus(397))],
    "MAX_PAG_POR_JANELA": 85,
    # Janela COMERCIAL (filtro por data de término da vigência):
    "VENCIDO_RECENTE_DIAS": 120,   # vencidos há até X dias ainda são "quentes"
    "FUTURO_MAX_DIAS": 180,        # vencendo em até X dias entram na lista
    "MONITOR_MAX_DIAS": 365,       # 181–365 dias só se alto valor / estratégico
    "VALOR_ALTO": 1_000_000,       # patamar de "alto valor"
    "OUT_DIR": core.CONFIG["OUT_DIR"],
    # Concorrentes reais (edite!): nomes/fragmentos -> marcados como Ameaça Alta
    "CONCORRENTES": [],            # ex.: ["oracle", "megasoft", "conexti"]
    "CARTEIRA_POR_UF": {"DF": "Conta-Chave DF", "GO": "Goiás"},
    "RESP_POR_UF": {"DF": "A definir (DF)", "GO": "A definir (GO)"},
}
core.CONFIG["CONCORRENTES_CONHECIDOS"] = ATK["CONCORRENTES"]  # o detector do core usa esta lista

ESTRATEGICAS = {"Cibersegurança", "Infraestrutura", "Cloud", "Dados/BI/IA"}

_LOG = []
def log(m):
    line = f"[{datetime.datetime.now():%H:%M:%S}] {m}"; _LOG.append(line); print(line, flush=True)

# =============================== COLETA (multi-janela) ===============================
def coletar():
    # Fetch PARALELO: a página 1 informa totalPaginas; as demais são baixadas
    # em pool de threads. DF/GO é extraído incrementalmente (memória limitada ao
    # subconjunto DF/GO, não às páginas nacionais). Cache em disco => resumível.
    from concurrent.futures import ThreadPoolExecutor, as_completed
    import datetime as _dt
    WORKERS = int(os.environ.get("ATLAS_FETCH_WORKERS", "3"))  # PNCP faz throttle acima de ~6 simultâneas
    MAXPAG = int(ATK["MAX_PAG_POR_JANELA"])
    MAX_DEPTH = 9          # profundidade máx. de subdivisão (365d -> ~1d)
    MIN_DIAS = 3           # não subdivide faixas menores que isso
    vistos, regs = set(), []
    def _absorve(j):
        for it in (j.get("data") or []):
            if (it.get("unidadeOrgao") or {}).get("ufSigla") in ATK["UFS"]:
                k = it.get("numeroControlePNCP") or f"{(it.get('orgaoEntidade') or {}).get('cnpj')}-{it.get('anoContrato')}-{it.get('sequencialContrato')}"
                if k not in vistos: vistos.add(k); regs.append(it)
    def _d(s):  return _dt.datetime.strptime(s, "%Y%m%d").date()
    def _s(d):  return d.strftime("%Y%m%d")
    def _meio(di, df):
        d0, d1 = _d(di), _d(df)
        m = d0 + (d1 - d0) // 2
        return _s(m), _s(m + _dt.timedelta(days=1))
    def _dias(di, df): return (_d(df) - _d(di)).days

    def _coleta(di, df, depth=0):
        """Coleta uma faixa de datas. Em caso de 500 cacheado na pág.1 OU de
        volume acima do teto de páginas, SUBDIVIDE a faixa ao meio e tenta as
        metades (faixa menor => assinatura de query diferente => escapa do 500;
        e nunca trunca silenciosamente uma janela grande)."""
        ini = len(regs)
        try:
            j1 = core.fetch_pagina(di, df, 1)
        except Exception as e:
            if depth < MAX_DEPTH and _dias(di, df) > MIN_DIAS:
                a, b = _meio(di, df)
                log(f"   ~ {di}-{df} pág1 falhou ({e}); subdividindo em {di}-{a} + {b}-{df}")
                _coleta(di, a, depth + 1); _coleta(b, df, depth + 1)
            else:
                log(f"   !! {di}-{df} pág1 falhou em definitivo ({e}) — faixa mínima atingida")
            return
        total = j1.get("totalRegistros"); tot_pags = int(j1.get("totalPaginas") or 1)
        # excede o teto -> subdivide (sem cap silencioso) em vez de descartar páginas
        if tot_pags > MAXPAG and depth < MAX_DEPTH and _dias(di, df) > MIN_DIAS:
            a, b = _meio(di, df)
            log(f"   ~ {di}-{df}: {tot_pags} págs > teto {MAXPAG}; subdividindo em {di}-{a} + {b}-{df}")
            _coleta(di, a, depth + 1); _coleta(b, df, depth + 1)
            return
        npags = min(tot_pags, MAXPAG)
        log(f"   janela {di}-{df}: {total} nacionais | {tot_pags} págs | coletando {npags} (workers={WORKERS})")
        _absorve(j1)
        if npags >= 2:
            def _fetch(p):
                try: return p, core.fetch_pagina(di, df, p)
                except Exception: return p, None
            falhas = []
            with ThreadPoolExecutor(max_workers=WORKERS) as ex:
                futs = [ex.submit(_fetch, p) for p in range(2, npags + 1)]
                done = 0
                for fut in as_completed(futs):
                    p, j = fut.result()
                    if j is not None: _absorve(j)
                    else: falhas.append(p)
                    done += 1
                    if done % 200 == 0: log(f"      {di}-{df}: {done}/{npags-1} págs | falhas {len(falhas)} | DF/GO acum {len(regs)}")
            # 2ª passada: retry SEQUENCIAL (conc=1) das páginas que falharam — amigável a throttle
            rodada = 0
            while falhas and rodada < 3:
                rodada += 1
                log(f"   retry seq (rodada {rodada}) de {len(falhas)} págs em {di}-{df}...")
                resto = []
                for p in sorted(falhas):
                    try: _absorve(core.fetch_pagina(di, df, p))
                    except Exception: resto.append(p)
                    time.sleep(core.CONFIG["SLEEP"])
                falhas = resto
            if falhas:
                log(f"   !! {di}-{df}: {len(falhas)} págs ainda falharam (re-rode p/ completar via cache): {sorted(falhas)[:15]}")
        log(f"   -> janela {di}-{df}: +{len(regs)-ini} registros DF/GO (total {len(regs)})")

    log(f"MODO ATAQUE | UFs {ATK['UFS']} | janelas {ATK['JANELAS']} | máx {MAXPAG} págs/janela | workers={WORKERS}")
    for (di, df) in ATK["JANELAS"]:
        _coleta(str(di), str(df), 0)
    log(f"Coleta concluída: {len(regs)} registros DF/GO únicos (em {len(ATK['JANELAS'])} janelas).")
    return regs

# =============================== LÓGICA COMERCIAL ===============================
def janela_comercial(dias, fim):
    if fim is None: return "Validação manual"
    if dias < 0:    return "Ataque imediato" if dias >= -ATK["VENCIDO_RECENTE_DIAS"] else "Monitoramento"
    if dias <= 30:  return "Ataque imediato"
    if dias <= 60:  return "Abordagem prioritária"
    if dias <= 90:  return "Preparação comercial"
    if dias <= 180: return "Nutrição estratégica"
    return "Monitoramento"

def urgencia(dias, fim, valor, categoria, score, conf):
    if fim is None or conf == "Baixa": return "Revisão"
    estrat = (valor >= ATK["VALOR_ALTO"]) or (categoria in ESTRATEGICAS)
    if (dias < 0 or dias <= 30) and estrat: return "Crítica"
    if dias <= 60 or score >= 80:          return "Alta"
    if dias <= 90:                         return "Média"
    return "Baixa"

def tipo_oportunidade(dias, fim, eh_conc):
    if eh_conc:        return "Substituição de concorrente"
    if fim is None:    return "Relacionamento (validar dados)"
    if dias < 0:       return "Renovação / nova licitação"
    if dias <= 90:     return "Renovação antecipada"
    return "Relacionamento / expansão"

ARGUMENTOS = {
 "Cibersegurança": "Mapear o cenário atual de segurança, validar a renovação do contrato vigente e posicionar solução alternativa antes da nova contratação.",
 "Infraestrutura": "Avaliar o ciclo de renovação de equipamentos, a capacidade atual e a modernização com proposta competitiva.",
 "Cloud": "Identificar se o órgão pretende renovar o ambiente em nuvem, migrar cargas ou otimizar custos operacionais.",
 "Software": "Verificar a dependência do fornecedor atual, o modelo de licenciamento e a possibilidade de substituição ou expansão.",
 "Serviços/Outsourcing": "Entender o nível de satisfação com o fornecedor atual e antecipar a abordagem para o novo ciclo de contratação.",
 "Dados/BI/IA": "Mapear a maturidade analítica do órgão e propor solução de BI, dados ou automação com base no contrato identificado.",
 "Redes": "Avaliar a infraestrutura de rede atual e posicionar modernização e segurança antes do próximo ciclo.",
 "Telecom": "Mapear o contrato de telecom vigente e antecipar proposta competitiva para a renovação.",
 "Hardware": "Identificar o ciclo de reposição de equipamentos e antecipar proposta para a próxima aquisição.",
 "CFTV/Videomonitoramento": "Avaliar o parque de videomonitoramento atual e propor modernização com analítico de vídeo.",
 "Backup/DR": "Validar a estratégia de backup e continuidade do órgão e posicionar solução de maior resiliência.",
 "Governança de TI": "Mapear o plano diretor de TI e posicionar serviços de governança e gestão.",
 "TI (geral)": "Mapear o escopo de TI contratado, validar o ciclo de renovação e posicionar proposta competitiva.",
}
def argumento(cat): return ARGUMENTOS.get(cat, ARGUMENTOS["TI (geral)"])

def proxima_acao(jc, dias, fim, eh_conc):
    base = {
     "Ataque imediato": "Acionar o responsável comercial HOJE; verificar se houve aditivo/renovação após o vencimento e se o órgão já está na carteira.",
     "Abordagem prioritária": "Pesquisar os decisores da unidade compradora e preparar a abordagem em até 48h.",
     "Preparação comercial": "Mapear o edital anterior e os requisitos técnicos; iniciar relacionamento com a área de TI do órgão.",
     "Nutrição estratégica": "Registrar relacionamento e revisar em 30 dias; preparar material técnico e prova de conceito.",
     "Monitoramento": "Marcar como monitoramento e revisar em 30 dias.",
     "Validação manual": "Encaminhar para a coordenação validar os dados e definir o responsável comercial.",
    }.get(jc, "Encaminhar para a coordenação.")
    if eh_conc: base += " Fornecedor atual é possível concorrente — preparar diferenciação técnica."
    return base

def score_ataque(dias, fim, valor, categoria, eh_conc, conf, esfera, uf, qual_frac):
    p = {}
    if fim is None:        p["venc"] = 0
    elif dias < 0:         p["venc"] = 24 if dias >= -ATK["VENCIDO_RECENTE_DIAS"] else 6
    elif dias <= 30:       p["venc"] = 30
    elif dias <= 60:       p["venc"] = 22
    elif dias <= 90:       p["venc"] = 15
    elif dias <= 180:      p["venc"] = 8
    else:                  p["venc"] = 3
    if   valor >= 5_000_000: p["valor"] = 20
    elif valor >= 2_000_000: p["valor"] = 17
    elif valor >= 1_000_000: p["valor"] = 14
    elif valor >=   500_000: p["valor"] = 10
    elif valor >=   100_000: p["valor"] = 6
    elif valor > 0:          p["valor"] = 3
    else:                    p["valor"] = 0
    p["categoria"]   = 15 if categoria in ESTRATEGICAS else (9 if categoria else 0)
    p["concorrente"] = 10 if eh_conc else 0
    p["confianca"]   = {"Alta": 8, "Média": 5, "Baixa": 2}.get(conf, 0)
    p["qualidade"]   = round(7 * qual_frac)
    p["orgao"]       = 5 if esfera in ("Federal", "Estadual", "Distrital") else (2 if esfera else 3)
    p["uf"]          = 3 if uf in ATK["UFS"] else 0
    p["fonte"]       = 2
    return min(100, sum(p.values())), p

def faixa(score):
    if score >= 90: return "Ataque máximo"
    if score >= 80: return "Prioridade crítica"
    if score >= 70: return "Alta prioridade"
    if score >= 50: return "Média prioridade"
    if score >= 30: return "Monitoramento"
    return "Baixa / revisão"

def explica_score(score, dias, fim, valor, categoria, eh_conc, conf):
    fs = []
    if fim is None: fs.append("sem data de término")
    elif dias < 0:  fs.append(f"contrato vencido há {abs(dias)} dias")
    else:           fs.append(f"vence em {dias} dias")
    if valor >= ATK["VALOR_ALTO"]: fs.append(f"valor de {core.brl(valor)}")
    if categoria in ESTRATEGICAS:  fs.append(f"categoria estratégica ({categoria})")
    if eh_conc:                    fs.append("fornecedor atual é possível concorrente")
    if conf == "Alta":             fs.append("classificação de TI de alta confiança")
    return f"Score {score}/100: " + "; ".join(fs) + "."

def venc_legivel(dias, fim):
    if fim is None: return "sem data"
    if dias < 0:    return f"venceu há {abs(dias)}d ({core.fmt_data(fim)})"
    return f"vence em {dias}d ({core.fmt_data(fim)})"

# =============================== PROCESSAMENTO ===============================
def processar(regs):
    log("Processando, classificando e priorizando (foco em janela comercial)...")
    # recorrência por fornecedor (nº de contratos na base coletada)
    rec_forn = collections.Counter((r.get("nomeRazaoSocialFornecedor") or "").strip() for r in regs)
    linhas = []
    for i, it in enumerate(regs, 1):
        org = it.get("orgaoEntidade") or {}; uni = it.get("unidadeOrgao") or {}
        objeto = it.get("objetoContrato") or ""
        cl = core.classifica_ti(objeto)
        if cl["eh_ti"] != "Sim":
            continue
        fim = core.parse_data(it.get("dataVigenciaFim"))
        dias = (fim - HOJE).days if fim else None
        # FILTRO DE JANELA COMERCIAL
        if fim is not None:
            if dias < -ATK["VENCIDO_RECENTE_DIAS"]:          # vencido há muito tempo -> fora
                continue
            if dias > ATK["FUTURO_MAX_DIAS"]:
                valor_tmp = core.to_float(it.get("valorGlobal")) or core.to_float(it.get("valorInicial"))
                estrat = (valor_tmp >= ATK["VALOR_ALTO"]) or (cl["categoria"] in ESTRATEGICAS)
                if not (dias <= ATK["MONITOR_MAX_DIAS"] and estrat):
                    continue
        valor = core.to_float(it.get("valorGlobal")) or core.to_float(it.get("valorInicial"))
        inicio = core.parse_data(it.get("dataVigenciaInicio")); assinatura = core.parse_data(it.get("dataAssinatura"))
        uf = uni.get("ufSigla") or ""; esfera = core.ESFERA.get(org.get("esferaId"), ""); poder = core.PODER.get(org.get("poderId"), "")
        forn = (it.get("nomeRazaoSocialFornecedor") or "").strip()
        eh_conc, grau = core.eh_concorrente(forn, "Sim")
        checks = [bool(core.limpa_cnpj(org.get("cnpj"))), bool(forn), valor > 0, fim is not None, bool(objeto.strip())]
        qfrac = sum(checks) / len(checks)
        qualidade = "Alta" if qfrac >= 0.8 else ("Média" if qfrac >= 0.6 else "Baixa")
        status, _ = core.status_contrato(fim)
        score, _p = score_ataque(dias, fim, valor, cl["categoria"], eh_conc, cl["confianca"], esfera, uf, qfrac)
        prio = faixa(score); jc = janela_comercial(dias, fim); urg = urgencia(dias, fim, valor, cl["categoria"], score, cl["confianca"])
        cnpj_org = core.limpa_cnpj(org.get("cnpj")); ano = it.get("anoContrato"); seq = it.get("sequencialContrato")
        link = f"https://pncp.gov.br/app/contratos/{cnpj_org}/{ano}/{seq}" if (cnpj_org and ano and seq) else "https://pncp.gov.br"
        nome_org = org.get("razaoSocial") or uni.get("nomeUnidade") or ""
        linhas.append({
            "ID Oportunidade": f"ATK-{i:05d}", "ID PNCP": it.get("numeroControlePNCP") or "", "Fonte": "PNCP",
            "Data da Coleta": HOJE.strftime("%d/%m/%Y"), "Link da Fonte": link,
            "Número do Contrato": it.get("numeroContratoEmpenho") or "", "Número do Processo": it.get("processo") or "",
            "Órgão": nome_org, "Nome Padronizado (Órgão)": core.pad_nome(nome_org), "CNPJ do Órgão": core.formata_cnpj(org.get("cnpj")),
            "UF": uf, "Município": uni.get("municipioNome") or "", "Poder": poder, "Esfera": esfera,
            "Unidade Compradora": uni.get("nomeUnidade") or "", "Segmento Presumido": core.segmento_presumido(nome_org),
            "Objeto": objeto.strip(), "Objeto Normalizado": core.norm(objeto)[:300],
            "Categoria Principal": cl["categoria"], "Subcategoria": cl["subcategoria"],
            "Palavras-chave Encontradas": ", ".join(cl["palavras"][:8]),
            "Valor Total": round(valor, 2),
            "Valor Mensal Estimado": round(valor / max(1, core.meses(inicio, fim)), 2) if (valor and inicio and fim) else "",
            "Data de Assinatura": core.fmt_data(assinatura), "Início da Vigência": core.fmt_data(inicio), "Fim da Vigência": core.fmt_data(fim),
            "Dias até Vencimento": dias if dias is not None else "", "Status do Contrato": status, "Janela Comercial": jc,
            "Fornecedor": forn, "Nome Padronizado (Fornecedor)": core.pad_nome(forn), "CNPJ do Fornecedor": core.formata_cnpj(it.get("niFornecedor")),
            "Possível Concorrente": "Sim" if eh_conc else "Não", "Grau de Ameaça": grau,
            "Recorrência do Fornecedor": rec_forn.get(forn, 1),
            "É TI?": "Sim", "Confiança da Classificação": cl["confianca"],
            "Score de Oportunidade": score, "Prioridade": prio, "Motivo da Prioridade": explica_score(score, dias, fim, valor, cl["categoria"], eh_conc, cl["confianca"]),
            "Tipo de Oportunidade": tipo_oportunidade(dias, fim, eh_conc),
            "Próxima Ação Recomendada": proxima_acao(jc, dias, fim, eh_conc),
            "Argumento Comercial Sugerido": argumento(cl["categoria"]),
            "Urgência Comercial": urg,
            "Responsável Comercial Sugerido": ATK["RESP_POR_UF"].get(uf, ""),
            "Carteira Sugerida": ATK["CARTEIRA_POR_UF"].get(uf, ""),
            "Necessita Revisão?": cl["revisao"], "Observações": "",
            "explicacao_score": explica_score(score, dias, fim, valor, cl["categoria"], eh_conc, cl["confianca"]),
        })
    # ordena por score desc, valor desc
    linhas.sort(key=lambda x: (x["Score de Oportunidade"], x["Valor Total"]), reverse=True)
    for n, l in enumerate(linhas, 1): l["ID Oportunidade"] = f"ATK-{n:05d}"
    log(f"   oportunidades de ataque (TI na janela comercial): {len(linhas)}")
    return linhas

# =============================== SAÍDAS ===============================
def csv_out(path, linhas, cols):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore"); w.writeheader()
        for l in linhas: w.writerow(l)
    log(f"   -> {os.path.basename(path)} ({len(linhas)} linhas)")

PRIO_FILL = {"Ataque máximo": "7F1D1D", "Prioridade crítica": "DC2626", "Alta prioridade": "D97706",
             "Média prioridade": "CA8A04", "Monitoramento": "2563EB", "Baixa / revisão": "64748B"}
URG_FILL  = {"Crítica": "DC2626", "Alta": "D97706", "Média": "CA8A04", "Baixa": "2563EB", "Revisão": "64748B"}

def xlsx_out(path, linhas, agg):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    HF = Font(bold=True, color="FFFFFF", name="Calibri"); HFILL = PatternFill("solid", fgColor="0F172A")
    thin = Side(style="thin", color="E2E8F0"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    def header(ws, row, cols, widths=None):
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row, i, c); cell.font = HF; cell.fill = HFILL; cell.alignment = Alignment(horizontal="left", vertical="center")
            if widths: ws.column_dimensions[get_column_letter(i)].width = widths[i-1]
        ws.row_dimensions[row].height = 22

    # --- Resumo Executivo
    ws = wb.active; ws.title = "Resumo Executivo"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1"); t = ws["A1"]; t.value = "ATLAS B2G — Lista de Ataque Comercial Imediato"
    t.font = Font(bold=True, size=16, color="0F172A"); t.alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:D2"); ws["A2"] = f"DF + Goiás · Tecnologia · gerado em {HOJE:%d/%m/%Y} · fonte PNCP (dados públicos)"
    ws["A2"].font = Font(italic=True, color="64748B")
    kpis = [("Oportunidades de ataque", agg["n"]), ("Críticas (Ataque máx./Crítica)", agg["criticas"]),
            ("Valor total mapeado", core.brl(agg["valor"])), ("Vencidos recentes", agg["vencidos"]),
            ("Vencendo ≤30 dias", agg["v30"]), ("Vencendo 31–90 dias", agg["v3190"]),
            ("DF", agg["uf"].get("DF", 0)), ("GO", agg["uf"].get("GO", 0))]
    header(ws, 4, ["Indicador", "Valor"], [34, 26, 12, 12])
    for r, (k, v) in enumerate(kpis, 5):
        ws.cell(r, 1, k).border = BORD; c = ws.cell(r, 2, v); c.border = BORD; c.font = Font(bold=True, color="2563EB")
    # Top 10 ao lado (sem reescrever A4/B4)
    for cc, txt in ((3, "Top 10 — Órgão"), (4, "Score")):
        cell = ws.cell(4, cc, txt); cell.font = HF; cell.fill = HFILL; cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.column_dimensions["C"].width = 46; ws.column_dimensions["D"].width = 10
    for r, l in enumerate(linhas[:10], 5):
        ws.cell(r, 3, f"{l['Nome Padronizado (Órgão)']} ({l['UF']}) · {l['Categoria Principal']}").border = BORD
        sc = ws.cell(r, 4, l["Score de Oportunidade"]); sc.border = BORD; sc.font = Font(bold=True)

    # --- Lista de Ataque (completa, badges, filtros)
    cols = list(linhas[0].keys()) if linhas else []
    cols = [c for c in cols if c != "explicacao_score"]
    wl = wb.create_sheet("Lista de Ataque")
    header(wl, 1, cols, [16]*len(cols))
    for r, l in enumerate(linhas, 2):
        for i, c in enumerate(cols, 1):
            cell = wl.cell(r, i, l.get(c, "")); cell.border = BORD; cell.alignment = Alignment(vertical="center", wrap_text=(c in ("Objeto","Motivo da Prioridade","Próxima Ação Recomendada","Argumento Comercial Sugerido")))
            if c == "Prioridade":
                cell.fill = PatternFill("solid", fgColor=PRIO_FILL.get(l[c], "FFFFFF")); cell.font = Font(bold=True, color="FFFFFF")
            elif c == "Urgência Comercial":
                cell.fill = PatternFill("solid", fgColor=URG_FILL.get(l[c], "FFFFFF")); cell.font = Font(bold=True, color="FFFFFF")
            elif c == "Valor Total":
                cell.number_format = '"R$" #,##0'
    # larguras úteis
    wide = {"Órgão":34, "Nome Padronizado (Órgão)":30, "Objeto":50, "Motivo da Prioridade":46,
            "Próxima Ação Recomendada":46, "Argumento Comercial Sugerido":48, "Fornecedor":28, "Categoria Principal":18}
    for i, c in enumerate(cols, 1):
        if c in wide: wl.column_dimensions[get_column_letter(i)].width = wide[c]
    wl.freeze_panes = "A2"; wl.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(linhas)+1}"

    # --- Top 10 Semana
    wt = wb.create_sheet("Top 10 Semana")
    tcols = ["Score de Oportunidade","Prioridade","Urgência Comercial","Nome Padronizado (Órgão)","UF","Categoria Principal","Valor Total","Status do Contrato","Dias até Vencimento","Fornecedor","Próxima Ação Recomendada"]
    header(wt, 1, tcols, [10,16,12,30,6,16,14,16,10,26,46])
    for r, l in enumerate(linhas[:10], 2):
        for i, c in enumerate(tcols, 1):
            cell = wt.cell(r, i, l.get(c, "")); cell.border = BORD; cell.alignment = Alignment(vertical="center", wrap_text=(c=="Próxima Ação Recomendada"))
            if c == "Prioridade": cell.fill = PatternFill("solid", fgColor=PRIO_FILL.get(l[c],"FFFFFF")); cell.font = Font(bold=True, color="FFFFFF")
            if c == "Valor Total": cell.number_format = '"R$" #,##0'
    wt.freeze_panes = "A2"

    # --- Por UF / Por Categoria / Por Fornecedor / Revisão / Parâmetros
    def aba_cont(nome, counter, c1, valor_map=None):
        w = wb.create_sheet(nome)
        if valor_map: header(w, 1, [c1, "Qtd", "Valor total"], [44, 10, 18])
        else: header(w, 1, [c1, "Qtd"], [44, 10])
        for r, (k, v) in enumerate(counter.most_common(60), 2):
            w.cell(r, 1, k or "—").border = BORD; w.cell(r, 2, v).border = BORD
            if valor_map:
                vc = w.cell(r, 3, round(valor_map.get(k, 0))); vc.number_format = '"R$" #,##0'; vc.border = BORD
    aba_cont("Por UF", agg["uf"], "UF", agg["uf_valor"])
    aba_cont("Por Categoria", agg["cat"], "Categoria", agg["cat_valor"])
    aba_cont("Por Fornecedor", agg["forn"], "Fornecedor", agg["forn_valor"])
    # Revisão Manual
    wr = wb.create_sheet("Revisão Manual")
    rcols = ["ID Oportunidade","Órgão","UF","Valor Total","Confiança da Classificação","Necessita Revisão?","Fim da Vigência","Objeto"]
    header(wr, 1, rcols, [14,34,6,14,14,14,14,60])
    rev = [l for l in linhas if l["Necessita Revisão?"] == "Sim" or l["Confiança da Classificação"] == "Baixa" or l["Fim da Vigência"] == ""]
    for r, l in enumerate(rev, 2):
        for i, c in enumerate(rcols, 1):
            cell = wr.cell(r, i, l.get(c, "")); cell.border = BORD; cell.alignment = Alignment(wrap_text=(c=="Objeto"))
            if c == "Valor Total": cell.number_format = '"R$" #,##0'
    wr.freeze_panes = "A2"
    # Parâmetros
    wp = wb.create_sheet("Parâmetros da Execução")
    header(wp, 1, ["Parâmetro", "Valor"], [34, 70])
    params = [("UFs", ", ".join(ATK["UFS"])), ("Janelas de publicação", " | ".join(f"{a}-{b}" for a,b in ATK["JANELAS"])),
        ("Máx páginas/janela", ATK["MAX_PAG_POR_JANELA"]), ("Vencido recente (dias)", ATK["VENCIDO_RECENTE_DIAS"]),
        ("Futuro máx (dias)", ATK["FUTURO_MAX_DIAS"]), ("Monitoramento máx (dias)", ATK["MONITOR_MAX_DIAS"]),
        ("Valor alto (R$)", ATK["VALOR_ALTO"]), ("Concorrentes configurados", ", ".join(ATK["CONCORRENTES"]) or "(nenhum — preencher)"),
        ("Categorias de TI (dicionário)", str(sum(len(v) for v in core.KW.values())) + " termos"),
        ("Data/hora da execução", f"{datetime.datetime.now():%d/%m/%Y %H:%M}")]
    for r, (k, v) in enumerate(params, 2):
        wp.cell(r, 1, k).border = BORD; wp.cell(r, 2, v).border = BORD
    wb.save(path); log(f"   -> {os.path.basename(path)}")

def vendedores_csv(path, linhas):
    cols = ["Prioridade","Urgência Comercial","Órgão","UF","Categoria","Valor","Situação do Contrato","Fornecedor Atual","Argumento Comercial","Próxima Ação","Responsável Sugerido","Link"]
    out = []
    for l in linhas:
        out.append({
            "Prioridade": l["Prioridade"], "Urgência Comercial": l["Urgência Comercial"],
            "Órgão": l["Nome Padronizado (Órgão)"], "UF": l["UF"], "Categoria": l["Categoria Principal"],
            "Valor": core.brl(l["Valor Total"]),
            "Situação do Contrato": venc_legivel(l["Dias até Vencimento"] if isinstance(l["Dias até Vencimento"], int) else None, _parse_fim(l["Fim da Vigência"])),
            "Fornecedor Atual": l["Fornecedor"] + (" (possível concorrente)" if l["Possível Concorrente"]=="Sim" else ""),
            "Argumento Comercial": l["Argumento Comercial Sugerido"], "Próxima Ação": l["Próxima Ação Recomendada"],
            "Responsável Sugerido": l["Responsável Comercial Sugerido"], "Link": l["Link da Fonte"],
        })
    csv_out(path, out, cols)

def _parse_fim(s):
    return core.parse_data(s) if s else None

# =============================== RELATÓRIO .MD ===============================
def relatorio_md(path, linhas, agg, total_regs, total_nac, dt):
    L = linhas
    topN = lambda key, n=10: sorted(L, key=lambda x: x[key] if isinstance(x[key], (int, float)) else 0, reverse=True)[:n]
    def linha_op(l): return f"- **{l['Nome Padronizado (Órgão)']}** ({l['UF']}) · {l['Categoria Principal']} · {core.brl(l['Valor Total'])} · {venc_legivel(l['Dias até Vencimento'] if isinstance(l['Dias até Vencimento'],int) else None,_parse_fim(l['Fim da Vigência']))} · _{l['Prioridade']}_ — {l['Próxima Ação Recomendada']}"
    md = []
    md.append("# ATLAS B2G — Lista de Ataque Comercial Imediato\n")
    md.append(f"**DF + Goiás · Tecnologia · fonte PNCP (dados públicos) · gerado em {HOJE:%d/%m/%Y}**\n")
    md.append("> Fila priorizada de oportunidades reais para a equipe comercial responder: *“quais órgãos atacar primeiro esta semana?”*\n")
    md.append("## 1. Resumo da coleta")
    md.append(f"- Janelas de publicação consultadas: {', '.join(f'{a}–{b}' for a,b in ATK['JANELAS'])}")
    md.append(f"- UFs: {', '.join(ATK['UFS'])}")
    md.append(f"- Registros DF/GO coletados: **{total_regs}** (de ~{total_nac} nacionais nas janelas)")
    md.append(f"- Oportunidades de TI na janela comercial: **{agg['n']}** · críticas: **{agg['criticas']}**")
    md.append(f"- Valor total mapeado: **{core.brl(agg['valor'])}**")
    md.append(f"- Vencidos recentes: **{agg['vencidos']}** · vencendo ≤30d: **{agg['v30']}** · 31–90d: **{agg['v3190']}**")
    md.append(f"- Tempo de execução: **{dt:.0f}s**\n")
    md.append("## 2. Ranking de oportunidades (Top 10 por score)")
    for l in L[:10]: md.append(linha_op(l))
    md.append("\n### Maiores por valor")
    for l in topN("Valor Total", 5): md.append(linha_op(l))
    md.append("\n## 3. Visão por UF")
    for uf in ATK["UFS"]:
        md.append(f"- **{uf}**: {agg['uf'].get(uf,0)} oportunidades · {core.brl(agg['uf_valor'].get(uf,0))} · categorias top: "
                  + ", ".join(f"{k}({v})" for k,v in agg['uf_cat'][uf].most_common(3)))
    md.append("\n## 4. Visão por categoria")
    for k, v in agg["cat"].most_common(12):
        md.append(f"- {k}: {v} oportunidades · {core.brl(agg['cat_valor'].get(k,0))}")
    md.append("\n## 5. Visão por fornecedor / concorrente")
    md.append("Fornecedores mais recorrentes (nº de contratos na base):")
    for k, v in agg["forn"].most_common(8):
        flag = " — possível concorrente" if k in agg["concorrentes"] else ""
        md.append(f"- {k}: {v} contrato(s) · {core.brl(agg['forn_valor'].get(k,0))}{flag}")
    md.append("\n## 6. Lista de ataque da semana (Top 10)")
    md.append("| # | Órgão | UF | Categoria | Valor | Situação | Prioridade | Ação |")
    md.append("|---|---|---|---|---|---|---|---|")
    for i, l in enumerate(L[:10], 1):
        sit = venc_legivel(l['Dias até Vencimento'] if isinstance(l['Dias até Vencimento'],int) else None,_parse_fim(l['Fim da Vigência']))
        md.append(f"| {i} | {l['Nome Padronizado (Órgão)']} | {l['UF']} | {l['Categoria Principal']} | {core.brl(l['Valor Total'])} | {sit} | {l['Prioridade']} | {l['Próxima Ação Recomendada'][:60]} |")
    md.append("\n## 7. Pendências de revisão")
    rev = [l for l in L if l["Necessita Revisão?"]=="Sim" or l["Confiança da Classificação"]=="Baixa" or l["Fim da Vigência"]==""]
    md.append(f"- Total que pedem curadoria: **{len(rev)}** (objeto ambíguo, baixa confiança ou sem data de término).")
    semfim = [l for l in L if l["Fim da Vigência"]==""]
    md.append(f"- Sem data de término: {len(semfim)} · Alto valor sem fornecedor claro: {len([l for l in L if l['Valor Total']>=ATK['VALOR_ALTO'] and not l['Fornecedor']])}")
    md.append("\n## 8. Recomendações")
    md.append("- **Distribuir primeiro** as oportunidades com prioridade *Ataque máximo / Crítica* e urgência *Crítica/Alta*.")
    md.append("- **Validar manualmente** os itens da aba *Revisão Manual* (objeto ambíguo / sem data).")
    md.append("- **Cadastrar concorrentes oficiais** em `ATK['CONCORRENTES']` (ex.: fornecedores recorrentes acima) — eleva o score e marca *Ameaça Alta*.")
    md.append("- **Foco de categoria**: priorizar as estratégicas (Cyber, Infra, Cloud, Dados/IA).")
    md.append("- **Ajuste de score/janela**: edite o bloco `ATK` (janelas, pesos, valor alto) para calibrar a fila.")
    md.append("- **Carregar na planilha**: importe `atlas_lista_ataque_comercial.csv` na aba *Oportunidades*; **no site**: converta o CSV para os dados do protótipo.")
    md.append("\n---\n*Gerado pelo ATLAS B2G — máquina de geração de oportunidades comerciais (piloto).*")
    with open(path, "w", encoding="utf-8") as f: f.write("\n".join(md))
    log(f"   -> {os.path.basename(path)}")

# =============================== AGREGADOS ===============================
def agregar(L):
    uf = collections.Counter(l["UF"] for l in L)
    cat = collections.Counter(l["Categoria Principal"] for l in L)
    forn = collections.Counter(l["Fornecedor"] for l in L if l["Fornecedor"])
    uf_valor = collections.defaultdict(float); cat_valor = collections.defaultdict(float); forn_valor = collections.defaultdict(float)
    uf_cat = {u: collections.Counter() for u in ATK["UFS"]}
    for l in L:
        uf_valor[l["UF"]] += l["Valor Total"]; cat_valor[l["Categoria Principal"]] += l["Valor Total"]
        if l["Fornecedor"]: forn_valor[l["Fornecedor"]] += l["Valor Total"]
        if l["UF"] in uf_cat: uf_cat[l["UF"]][l["Categoria Principal"]] += 1
    crit = sum(1 for l in L if l["Prioridade"] in ("Ataque máximo", "Prioridade crítica") or l["Urgência Comercial"] == "Crítica")
    venc = sum(1 for l in L if isinstance(l["Dias até Vencimento"], int) and l["Dias até Vencimento"] < 0)
    v30 = sum(1 for l in L if isinstance(l["Dias até Vencimento"], int) and 0 <= l["Dias até Vencimento"] <= 30)
    v3190 = sum(1 for l in L if isinstance(l["Dias até Vencimento"], int) and 31 <= l["Dias até Vencimento"] <= 90)
    concs = set(l["Fornecedor"] for l in L if l["Possível Concorrente"] == "Sim")
    return dict(n=len(L), valor=sum(l["Valor Total"] for l in L), criticas=crit, vencidos=venc, v30=v30, v3190=v3190,
                uf=uf, cat=cat, forn=forn, uf_valor=uf_valor, cat_valor=cat_valor, forn_valor=forn_valor,
                uf_cat=uf_cat, concorrentes=concs)

# =============================== MAIN ===============================
def main():
    t0 = time.time(); od = ATK["OUT_DIR"]; os.makedirs(od, exist_ok=True)
    log("ATLAS B2G — Lista de Ataque Comercial Imediato (início).")
    regs = coletar()
    total_nac = "—"
    L = processar(regs)
    if not L:
        log("Nenhuma oportunidade na janela comercial. Ajuste ATK['JANELAS'] / filtros."); _salva_log(od); return
    agg = agregar(L)
    cols = [c for c in L[0].keys() if c != "explicacao_score"]
    csv_out(os.path.join(od, "atlas_lista_ataque_comercial.csv"), L, cols + ["explicacao_score"])
    vendedores_csv(os.path.join(od, "atlas_top_oportunidades_para_vendedores.csv"), L)
    try: xlsx_out(os.path.join(od, "atlas_lista_ataque_comercial.xlsx"), L, agg)
    except Exception as e: log(f"   (xlsx falhou: {e})")
    relatorio_md(os.path.join(od, "atlas_relatorio_ataque_comercial.md"), L, agg, len(regs), total_nac, time.time()-t0)
    # resumo no log
    log("="*60); log(f"OPORTUNIDADES: {agg['n']} | CRÍTICAS: {agg['criticas']} | VALOR: {core.brl(agg['valor'])}")
    log(f"Vencidos: {agg['vencidos']} | ≤30d: {agg['v30']} | 31–90d: {agg['v3190']} | UF: {dict(agg['uf'])}")
    log("TOP 5:")
    for l in L[:5]:
        log(f"   [{l['Score de Oportunidade']}|{l['Prioridade']}] {l['Nome Padronizado (Órgão)']} ({l['UF']}) · {l['Categoria Principal']} · {core.brl(l['Valor Total'])}")
    log("="*60); log(f"Concluído em {time.time()-t0:.1f}s.")
    _salva_log(od)

def _salva_log(od):
    with open(os.path.join(od, "atlas_pncp_log_execucao_ataque.txt"), "w", encoding="utf-8") as f:
        f.write("PARÂMETROS:\n" + "\n".join(f"  {k} = {v}" for k, v in ATK.items()) + "\n\nLOG:\n" + "\n".join(_LOG))
    log(f"   -> atlas_pncp_log_execucao_ataque.txt")

if __name__ == "__main__":
    main()
