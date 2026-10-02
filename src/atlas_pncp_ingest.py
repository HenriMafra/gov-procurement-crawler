# -*- coding: utf-8 -*-
"""
ATLAS B2G — Ingestão Real do PNCP (Piloto DF + Goiás + Tecnologia)
==================================================================
Transforma contratos públicos do PNCP em inteligência comercial:
coleta -> normaliza -> classifica (TI) -> aplica regras comerciais
-> calcula score/prioridade -> gera bases e relatório executivo.

Somente DADOS PÚBLICOS. Apenas requisições GET (leitura). Sem credenciais.

API: https://pncp.gov.br/api/consulta/v1/contratos
     params: dataInicial=YYYYMMDD & dataFinal=YYYYMMDD & pagina=N [& tamanhoPagina]
     (filtra por data de PUBLICAÇÃO; não há filtro server-side por UF -> filtramos no cliente)

Como rodar:   python atlas_pncp_ingest.py
Configuração: edite o bloco CONFIG abaixo (UFs, janela de datas, páginas, valor mínimo...).
Reexecutável: idempotente; sobrescreve as saídas na pasta OUT_DIR.
"""

import sys, os, json, csv, re, time, unicodedata, datetime, urllib.request, urllib.error, ssl, collections

# reconfigura stdout p/ UTF-8 (console Windows)
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass

# =============================== CONFIG ===============================
HOJE = datetime.date.today()
def _ymd(d): return d.strftime("%Y%m%d")

CONFIG = {
    "UFS": ["DF", "GO"],                 # UFs do piloto
    # Janela de PUBLICAÇÃO no PNCP. Padrão: ~24 meses atrás (capta contratos
    # de 12–24 meses cuja vigência termina por volta de agora = oportunidades).
    "DATA_INICIAL": _ymd(HOJE - datetime.timedelta(days=744)),
    "DATA_FINAL":   _ymd(HOJE - datetime.timedelta(days=730)),
    "MAX_PAGINAS": 90,                   # teto de páginas (cada página = até 500)
    "TAMANHO_PAGINA": 500,               # padrão da API
    "VALOR_MIN": 0,                      # valor mínimo do contrato p/ entrar (0 = todos)
    "OPP_VALOR_MIN": 50000,              # valor mínimo p/ virar oportunidade
    "OPP_DIAS_MAX": 180,                 # vence em até X dias -> oportunidade
    "APENAS_TI": True,                   # base tratada só com contratos de TI
    "SLEEP": 0.35,                       # pausa educada entre requisições (s)
    "RETRIES": 4,                        # tentativas por página
    "TIMEOUT": 45,                       # timeout por requisição (s)
    "OUT_DIR": r"C:\Users\henri\Desktop\ATLAS-PNCP-Pilot",
    "USAR_CACHE": True,                  # cacheia páginas baixadas (re-execução instantânea)
    # Concorrentes conhecidos (preencha com fragmentos do nome dos SEUS rivais reais):
    "CONCORRENTES_CONHECIDOS": [],       # ex.: ["rivalsec", "infranuvem"]
    # Carteira/responsável sugeridos por UF (edite conforme seu time):
    "CARTEIRA_POR_UF": {"DF": "Conta-Chave DF", "GO": "Goiás"},
    "RESP_POR_UF":     {"DF": "A definir (DF)", "GO": "A definir (GO)"},
}
BASE_URL = "https://pncp.gov.br/api/consulta/v1/contratos"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) ATLAS-B2G-Pilot/1.0",
      "Accept": "application/json"}
_SSL = ssl.create_default_context()

# log em memória + arquivo
_LOG = []
def log(msg):
    ts = datetime.datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    _LOG.append(line); print(line, flush=True)

# =============================== HTTP ===============================
def _request(url):
    """GET com retry exponencial e timeout. Retorna dict JSON ou levanta."""
    last = None
    for tent in range(1, CONFIG["RETRIES"] + 1):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=CONFIG["TIMEOUT"], context=_SSL) as r:
                return json.loads(r.read().decode("utf-8", "replace"))
        except urllib.error.HTTPError as e:
            last = f"HTTP {e.code}"
            if e.code == 204:        # sem conteúdo
                return {"data": [], "totalPaginas": 0, "paginasRestantes": 0}
            if e.code in (400, 404): # erro definitivo de parâmetro
                raise
        except Exception as e:
            last = f"{type(e).__name__}: {str(e)[:80]}"
        espera = CONFIG["SLEEP"] * (2 ** (tent - 1))
        log(f"   tentativa {tent} falhou ({last}); aguardando {espera:.1f}s")
        time.sleep(espera)
    raise RuntimeError(f"Falha após {CONFIG['RETRIES']} tentativas: {last}")

def fetch_pagina(di, df, pagina):
    url = f"{BASE_URL}?dataInicial={di}&dataFinal={df}&pagina={pagina}"
    if CONFIG.get("USAR_CACHE"):
        cdir = os.path.join(CONFIG["OUT_DIR"], ".cache_pncp")
        os.makedirs(cdir, exist_ok=True)
        cf = os.path.join(cdir, f"{di}_{df}_{pagina}.json")
        if os.path.exists(cf):
            try: return json.load(open(cf, encoding="utf-8"))
            except Exception: pass
        j = _request(url)
        try: json.dump(j, open(cf, "w", encoding="utf-8"))
        except Exception: pass
        return j
    return _request(url)

# =============================== NORMALIZAÇÃO ===============================
def strip_acentos(s):
    if not s: return ""
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")

def norm(s):
    """minúsculo, sem acento, espaços colapsados (para matching/comparação)."""
    return re.sub(r"\s+", " ", strip_acentos(s).lower()).strip()

def limpa_cnpj(s):
    return re.sub(r"\D", "", str(s or ""))

def formata_cnpj(s):
    d = limpa_cnpj(s)
    if len(d) == 14:
        return f"{d[0:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:14]}"
    return s or ""

def pad_nome(s):
    """Title Case enxuto preservando siglas; remove espaços duplicados."""
    s = re.sub(r"\s+", " ", (s or "").strip())
    if not s: return ""
    minus = {"de","da","do","das","dos","e","em","no","na","a","o","ao","aos","às","à"}
    # palavras REAIS que às vezes vêm em CAIXA ALTA e NÃO são sigla (não manter em maiúsculo)
    nao_sigla = {"meio","mais","casa","rede","novo","nova","real","alto","alta","alem","além","vida",
                 "luz","sul","rio","mar","sao","são","ouro","vale","povo","agua","água","pico","lago",
                 "boa","bom","dois","tres","três","leste","oeste","norte","ilha","sede","obra","obras","area","área"}
    out = []
    for w in s.split(" "):
        wl = w.lower()
        if wl in minus:                                          # conector: minúsculo
            out.append(wl)
        elif w.isupper() and len(w) <= 4 and wl not in nao_sigla:  # sigla curta (TI, DF, SEDF...) — mas não palavra real
            out.append(w)
        else:
            out.append(w.capitalize())
    r = " ".join(out)
    return r[0].upper() + r[1:] if r else r

def parse_data(s):
    if not s: return None
    s = str(s)[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try: return datetime.datetime.strptime(s, fmt).date()
        except Exception: pass
    return None

def to_float(v):
    try:
        if v is None or v == "": return 0.0
        return float(v)
    except Exception:
        return 0.0

def fmt_data(d): return d.strftime("%d/%m/%Y") if d else ""

# =============================== CLASSIFICADOR DE TI ===============================
# Termos FORTES (qualquer um => é TI). Mapeados por categoria.
KW = {
 "Cibersegurança": ["firewall","ngfw","waf","siem","soc","edr","xdr","antivirus","ciberseguranca",
    "seguranca da informacao","seguranca cibernetica","criptografia","pentest","teste de intrusao",
    "endpoint","vpn","dlp","iam","mfa","controle de acesso logico","gestao de vulnerabilidade","ips","ids",
    # >>> Portfólio Enterprise IT Group (ZivaSec/EnterpriseCore/InfoSec/CloudSec) - pesquisa 2026-07-01 <<<
    "ztna","zero trust","zero trust network access","mdr","pam","iam/pam","casb",
    "threat hunting","sandboxing","gestao de acesso privilegiado","acesso privilegiado",
    "cofre digital","cofre de credenciais","cyberark","varonis","checkmarx","qualys","wiz",
    "netskope","gigamon","vectra","trendai","trend micro","trend vision one","check point",
    "checkpoint","cohesity","tuvis","everpure","security center","siggo","threat intelligence",
    "resposta a incidentes","gestao de incidentes de seguranca","analise de codigo","sast","dast",
    "protecao de e-mail","email security","sandbox de e-mail","managed detection and response",
    "deteccao e resposta estendida","xdr flex","soc cognitivo","ngsx","cidades inteligentes","smart cities",
    # >>> conferência visual da pag. Parceiros (screenshot) - completando o que faltou <<<
    "f5","f5 networks","cloudflare","cymulate","elastic","elastic security","fortinet","fortigate",
    "veracode","talsec","simulacao de ataque","gestao de exposicao a ameacas","seguranca de aplicativos moveis",
    # >>> achados nas paginas Academy/Vagas (job postings reais) - 2026-07-01 <<<
    "aruba","aruba networking","f5 big-ip","big-ip","seguranca ofensiva","red team","blue team",
    "analise e resposta a incidentes","reducao de falsos positivos",
    # >>> achado no LinkedIn oficial do Enterprise IT Group - 2026-07-01 <<<
    "thales","cloud checker"],
 "Infraestrutura": ["storage","datacenter","data center","no-break","nobreak",
    "virtualizacao","hypervisor","hiperconverg","sala cofre","rack para servidor","disaster recovery",
    "hpe","hewlett packard enterprise","nutanix","infraestrutura de redes","monitoramento de ativos","noc"],
 "Redes": ["switch","switches","roteador","roteadores","access point","ponto de acesso","rede logica",
    "rede de dados","cabeamento estruturado","fibra optica","sd-wan","controladora wireless","ativos de rede","wi-fi","wifi",
    "cisco","cisco systems","aruba","aruba networking",
    "gigamon","vectra","visibilidade de trafego","micro segmentacao","orquestracao de rede"],
 "Cloud": ["nuvem","cloud","iaas","paas","saas","computacao em nuvem","microsoft azure","azure",
    "amazon web services","aws","google cloud","oracle cloud","hospedagem em nuvem","backup em nuvem",
    "netskope","casb","seguranca em nuvem","cloud security","migracao para nuvem",
    "cnapp","multicloud","protecao de workloads","cloud native application protection platform"],
 "Software": ["software","licenca de software","licenciamento de software","erp","banco de dados",
    "business intelligence","power bi","sistema de gestao","sistema informatizado","solucao de software",
    "aplicativo","plataforma digital"],
 "Backup/DR": ["backup","copia de seguranca","draas","recuperacao de desastre",
    "cohesity","netbackup","enterprise vault","bcs premier","backup na nuvem","backup unificado"],
 "Serviços/Outsourcing": ["service desk","help desk","suporte tecnico","sustentacao de sistemas",
    "fabrica de software","outsourcing de ti","outsourcing","desenvolvimento de sistemas",
    "desenvolvimento de software","alocacao de mao de obra de ti","consultoria em ti"],
 "CFTV/Videomonitoramento": ["cftv","circuito fechado de televisao","videomonitoramento","camera ip",
    "cameras de seguranca","nvr","reconhecimento facial"],
 "Telecom": ["telefonia ip","voip","pabx","central telefonica","telecomunicacoes","link dedicado","sip trunk"],
 "Dados/BI/IA": ["inteligencia artificial","machine learning","aprendizado de maquina","big data",
    "data lake","data warehouse","ciencia de dados","governanca de dados","analytics","etl"],
 "Hardware": ["microcomputador","computador","computadores","notebook","desktop","estacao de trabalho",
    "equipamento de informatica","equipamentos de informatica","scanner","tablet","impressora"],
 "Governança de TI": ["governanca de ti","plano diretor de tecnologia","pdti","cobit","itil"],
 "TI (geral)": ["tecnologia da informacao","informatica","solucao de ti","servico de informatica","servicos de informatica"],
}
# >>> termos adicionais (taxonomia comercial ENTERPRISECORE) - 2026-06 <<<
for _k in list(KW.keys()):
    _kl=_k.lower()
    if _kl.startswith('ciber') or 'seguran' in _kl:
        KW[_k]=KW[_k]+['pam', 'acesso privilegiado', 'acessos privilegiados', 'cofre de senhas', 'gestao de identidades', 'drsi', 'deteccao e resposta a incidentes', 'mdr', 'centro de operacoes de seguranca', 'gerenciamento de logs', 'seguranca de rede', 'firewall de aplicacao web', 'web application firewall', 'antispam', 'gestao de vulnerabilidades', 'protecao contra ameacas ciberneticas', 'seguranca em nuvem', 'casb', 'cspm', 'privacidade de dados', 'lgpd', 'gdpr', 'active directory', 'dados nao estruturados', 'controlador de dominio']
    elif _kl.startswith('rede'):
        KW[_k]=KW[_k]+['wireless', 'rede sem fio', 'redes sem fio', 'wlan', 'hotspot', 'comutador de rede']

# >>> termos adicionais lote2 ENTERPRISECORE - 2026-06 <<<
for _k in list(KW.keys()):
    _kl=_k.lower()
    if _kl.startswith('ciber') or 'seguran' in _kl:
        KW[_k]=KW[_k]+['protecao de dados', 'seguranca de perimetro', 'protecao de perimetro', 'defesa de perimetro', 'defesas de perimetro', 'monitoramento de perimetro', 'controle de acesso de perimetro', 'deteccao de intrusao', 'deteccao de intrusoes', 'seguranca de borda de rede', 'seguranca de borda', 'analise de seguranca', 'monitoramento de seguranca', 'gestao de incidentes de seguranca', 'gerenciamento de incidentes de seguranca', 'security center', 'gpo', 'politicas de grupo', 'gerenciamento de politicas de grupo', 'gravacao de sessao', 'privileged password management', 'gerenciamento de senhas privilegiadas', 'controle de acesso a dados', 'controle de acesso aos dados', 'controle de acesso a arquivos', 'monitoramento de alteracoes de arquivos', 'auditoria de arquivos', 'analise de metadados', 'auditoria de dados', 'epp', 'protecao de servidores', 'manutencao de hardware', 'gerenciamento de vulnerabilidades', 'analise de vulnerabilidade', 'compliance de dados', 'conformidade de dados', 'protecao de aplicativos web', 'seguranca de aplicativos web', 'seguranca de aplicacoes web', 'protecao contra ataques web', 'protecao contra ataques na web', 'lei de protecao de dados', 'politica de privacidade', 'gestao de riscos de dados', 'protecao contra violacao de dados', 'violacao de dados', 'recuperacao de desastres']
    elif _kl.startswith('rede'):
        KW[_k]=KW[_k]+['controladora wireless', 'ponto de acesso sem fio', 'conectividade sem fio']

# Subconjunto MUITO forte (1 ocorrência já dá confiança Alta)
MUITO_FORTE = {"firewall","ngfw","siem","edr","xdr","waf","ciberseguranca","seguranca da informacao",
 "datacenter","data center","servidor","servidores","switch","switches","roteador","erp","backup",
 "nuvem","cloud","cftv","videomonitoramento","software","banco de dados","service desk","help desk",
 "fabrica de software","desenvolvimento de sistemas","desenvolvimento de software","inteligencia artificial",
 "tecnologia da informacao","microcomputador","notebook","storage","virtualizacao",
 # >>> marcas/fabricantes do portfólio Enterprise IT Group - 2026-07-01 <<<
 "cyberark","varonis","checkmarx","qualys","wiz","netskope","gigamon","vectra","trendai",
 "check point","checkpoint","cohesity","ztna","zero trust","mdr","pam","casb","threat hunting"}
# Termos AMBÍGUOS: só contam se houver qualificador de TI por perto (senão ignorados)
AMBIGUOS = ["manutencao","suporte","rede","sistema","monitoramento","servicos","plataforma","licenca",
 "digital","automacao","integracao","tecnologia","solucao","hospedagem"]
QUALIF_TI = ["computador","informatica","de ti","de dados","de rede","software","sistema","servidor",
 "rede logica","tecnologia da informacao","nuvem"]
# Contexto NEGATIVO (se presente e SEM termo forte de TI -> NÃO é TI)
NEGATIVO = ["ar condicionado","climatizacao","automotivo","veiculo","veiculos","predial","obra","obras",
 "reforma","alimento","alimenticio","merenda","medicamento","hospitalar","saude","fonoaudiolog","odontolog",
 "eletricista","encanador","limpeza","jardinagem","pintura","mobiliario","movel","moveis","combustivel",
 "pneu","engenharia civil","transporte escolar","vigilancia armada","seguranca patrimonial","seguro veicular",
 "uniforme","genero alimenticio","material de construcao","poda","capina","funerari","marketing","grafica","grafico"]

# índice termo -> categoria
_TERMO_CAT = {}
for cat, termos in KW.items():
    for t in termos: _TERMO_CAT[t] = cat
_TERMO_CAT["servidor"] = "Infraestrutura"   # tratado em caso especial (ambíguo: servidor de TI x servidor público)
_TODOS_FORTES = list(_TERMO_CAT.keys())
# marcadores de "servidor público" (funcionário) -> NÃO é servidor de TI
_SERVIDOR_CIVIL = ["servidor publico","servidores publicos","do servidor","dos servidores","imuniza",
 "vacina","aposentad","capacitacao","pericia","folha de pagamento","saude do servidor","exame medico",
 "auxilio","gratificacao","estatutario","efetivo","psicolog","clinica","odontolog","atendimento medico",
 "beneficio","empenho em favor"]
# 'servidor' só é servidor de TI se acompanhado de um qualificador técnico:
_SERVIDOR_QUAL = ["rede","de dados","aplicacao","aplicacoes","blade","rack","virtual","banco de dados",
 "web","arquivos","corporativo","datacenter","data center","hospedagem","storage","de email","de e-mail",
 "torre","linux","windows server","fisico","de banco","de imagem","de arquivo","cpd"]

def _tem(termo, texto):
    """match por palavra inteira no texto normalizado."""
    return re.search(r"(?<![a-z0-9])" + re.escape(termo) + r"(?![a-z0-9])", texto) is not None

def classifica_ti(objeto):
    """
    Retorna dict: eh_ti, categoria, subcategoria, confianca, palavras (lista),
    necessita_revisao.
    """
    txt = norm(objeto)
    if not txt:
        return dict(eh_ti="Não", categoria="", subcategoria="", confianca="—", palavras=[], revisao="Sim")
    achados = [t for t in _TODOS_FORTES if t != "servidor" and _tem(t, txt)]
    # caso especial 'servidor' (ambíguo): só conta como TI se houver qualificador técnico e
    # nenhum marcador de servidor público/funcionário.
    if (_tem("servidor", txt) or _tem("servidores", txt)) and not any(m in txt for m in _SERVIDOR_CIVIL):
        if any(_tem(q, txt) for q in _SERVIDOR_QUAL):
            achados.append("servidor")
    cats = collections.Counter(_TERMO_CAT[t] for t in achados)
    neg = any(_tem(n, txt) for n in NEGATIVO)

    if achados:
        categoria = cats.most_common(1)[0][0]
        # subcategoria: termo mais "específico" achado dessa categoria (o mais longo)
        termos_cat = [t for t in achados if _TERMO_CAT[t] == categoria]
        subcat = max(termos_cat, key=len) if termos_cat else ""
        muito = any(t in MUITO_FORTE for t in achados)
        if len(achados) >= 2 or muito:
            conf = "Alta"
        else:
            conf = "Média"
        revisao = "Sim" if (neg or conf == "Média") else "Não"
        return dict(eh_ti="Sim", categoria=categoria, subcategoria=subcat,
                    confianca=conf, palavras=achados, revisao=revisao)

    # sem termo forte: tenta ambíguo + qualificador
    amb = [a for a in AMBIGUOS if _tem(a, txt)]
    if amb and not neg:
        qual = [q for q in QUALIF_TI if _tem(q, txt)]
        if qual:
            return dict(eh_ti="Sim", categoria="TI (geral)", subcategoria=amb[0],
                        confianca="Baixa", palavras=amb + qual, revisao="Sim")
    return dict(eh_ti="Não", categoria="", subcategoria="", confianca="—", palavras=[], revisao="Não")

# =============================== REGRAS COMERCIAIS ===============================
def dias_ate_venc(fim):
    return (fim - HOJE).days if fim else None

def status_contrato(fim):
    d = dias_ate_venc(fim)
    if fim is None: return "Sem data de término", d
    if d < 0:   return "Vencido", d
    if d <= 30: return "Vencendo em 30 dias", d
    if d <= 60: return "Vencendo em 60 dias", d
    if d <= 90: return "Vencendo em 90 dias", d
    return "Ativo", d

ESTRATEGICAS = {"Cibersegurança", "Infraestrutura", "Cloud", "Dados/BI/IA"}
def eh_concorrente(forn_nome, eh_ti):
    n = norm(forn_nome)
    for c in CONFIG["CONCORRENTES_CONHECIDOS"]:
        if c and norm(c) in n:
            return True, "Alto"
    if eh_ti == "Sim":
        tokens = ["tecnologia","sistemas","informatica","software","solucoes","solucao",
                  "telecom","redes","seguranca","dados","digital","consultoria"]
        if any(_tem(t, n) for t in tokens):
            return True, "Médio"   # possível concorrente (a validar)
    return False, "—"

def score_oportunidade(dias, status, valor, categoria, eh_conc, esfera, uf, qualidade_frac):
    """Score 0–100 com pesos (ver doc). Retorna (score, partes)."""
    p = {}
    # Vencimento: até 30
    if status == "Vencido": p["vencimento"] = 20
    elif dias is None:      p["vencimento"] = 5
    elif dias <= 30:        p["vencimento"] = 30
    elif dias <= 60:        p["vencimento"] = 22
    elif dias <= 90:        p["vencimento"] = 15
    elif dias <= 180:       p["vencimento"] = 8
    else:                   p["vencimento"] = 3
    # Valor: até 20
    if   valor >= 5_000_000: p["valor"] = 20
    elif valor >= 2_000_000: p["valor"] = 17
    elif valor >= 1_000_000: p["valor"] = 14
    elif valor >=   500_000: p["valor"] = 10
    elif valor >=   100_000: p["valor"] = 6
    elif valor > 0:          p["valor"] = 3
    else:                    p["valor"] = 0
    # Categoria estratégica: até 15
    p["categoria"] = 15 if categoria in ESTRATEGICAS else (10 if categoria else 0)
    # Concorrente: até 10
    p["concorrente"] = 10 if eh_conc else 0
    # Órgão prioritário (esfera): até 8
    p["orgao"] = 8 if esfera in ("Federal","Estadual","Distrital") else (4 if esfera else 6)
    # UF prioritária: até 5
    p["uf"] = 5 if uf in CONFIG["UFS"] else 0
    # Qualidade do dado: até 7
    p["qualidade"] = round(7 * qualidade_frac)
    # Fonte confiável (PNCP): até 5
    p["fonte"] = 5
    score = min(100, sum(p.values()))
    return score, p

def faixa_prioridade(score):
    if score >= 85: return "Máxima"
    if score >= 70: return "Alta"
    if score >= 50: return "Média"
    if score >= 30: return "Baixa"
    return "Monitoramento"

def motivo_prioridade(score, dias, status, valor, categoria, eh_conc):
    partes = []
    if status == "Vencido": partes.append("contrato vencido")
    elif dias is not None and dias <= 180: partes.append(f"vence em {dias} dias")
    if valor >= 1_000_000: partes.append(f"valor de {brl(valor)}")
    if categoria in ESTRATEGICAS: partes.append(f"categoria estratégica ({categoria})")
    if eh_conc: partes.append("fornecedor atual é possível concorrente")
    base = "; ".join(partes) if partes else "dados disponíveis para análise"
    return f"Score {score}/100: " + base + "."

def proxima_acao(eh_ti, status, dias):
    if eh_ti != "Sim": return "Arquivar ou manter apenas como referência"
    if status == "Sem data de término": return "Enviar para validação manual"
    if status == "Vencido": return "Verificar renovação, aditivo ou novo processo"
    if dias is None: return "Enviar para validação manual"
    if dias <= 30:  return "Acionar responsável comercial imediatamente"
    if dias <= 60:  return "Preparar abordagem e mapear decisores"
    if dias <= 90:  return "Iniciar relacionamento e validar cenário"
    if dias <= 180: return "Monitorar e planejar abordagem"
    return "Monitorar (vencimento distante)"

def brl(v):
    try: return "R$ " + format(float(v), ",.0f").replace(",", ".")
    except Exception: return "R$ 0"

ESFERA = {"F": "Federal", "E": "Estadual", "M": "Municipal", "D": "Distrital", "N": ""}
PODER  = {"E": "Executivo", "L": "Legislativo", "J": "Judiciário", "N": ""}

# =============================== COLETA ===============================
def coletar():
    di, df = CONFIG["DATA_INICIAL"], CONFIG["DATA_FINAL"]
    log(f"Coletando PNCP | janela publicação {di}–{df} | UFs {CONFIG['UFS']} | máx {CONFIG['MAX_PAGINAS']} págs")
    registros, total_nac, pagina = [], None, 1
    while pagina <= CONFIG["MAX_PAGINAS"]:
        try:
            j = fetch_pagina(di, df, pagina)
        except Exception as e:
            log(f"!! parando na página {pagina}: {e}"); break
        data = j.get("data") or []
        if total_nac is None:
            total_nac = j.get("totalRegistros")
            log(f"   totalRegistros nacional na janela: {total_nac} | totalPaginas: {j.get('totalPaginas')}")
        if not data:
            log("   página vazia — fim."); break
        na_uf = [it for it in data if (it.get("unidadeOrgao") or {}).get("ufSigla") in CONFIG["UFS"]]
        registros.extend(na_uf)
        log(f"   pág {pagina}: {len(data)} regs | DF/GO acumulado: {len(registros)}")
        if (j.get("paginasRestantes") or 0) <= 0: break
        pagina += 1
        time.sleep(CONFIG["SLEEP"])
    log(f"Coleta concluída: {len(registros)} registros DF/GO (de ~{total_nac} nacionais na janela).")
    return registros, total_nac

# =============================== PROCESSAMENTO ===============================
def construir_linha(it, idx):
    org = it.get("orgaoEntidade") or {}
    uni = it.get("unidadeOrgao") or {}
    objeto = it.get("objetoContrato") or ""
    uf = uni.get("ufSigla") or ""
    esfera = ESFERA.get(org.get("esferaId"), "")
    poder = PODER.get(org.get("poderId"), "")
    fim = parse_data(it.get("dataVigenciaFim"))
    inicio = parse_data(it.get("dataVigenciaInicio"))
    assinatura = parse_data(it.get("dataAssinatura"))
    valor = to_float(it.get("valorGlobal")) or to_float(it.get("valorInicial"))
    cl = classifica_ti(objeto)
    status, dias = status_contrato(fim)
    forn_nome = (it.get("nomeRazaoSocialFornecedor") or "").strip()
    eh_conc, grau = eh_concorrente(forn_nome, cl["eh_ti"])
    # qualidade do dado (fração 0-1)
    checks = [bool(limpa_cnpj(org.get("cnpj"))), bool(forn_nome), valor > 0, fim is not None, bool(objeto.strip())]
    qual_frac = sum(checks) / len(checks)
    qualidade = {1.0: "Alta", 0.8: "Alta"}.get(qual_frac, "Média" if qual_frac >= 0.6 else "Baixa")
    score, partes = score_oportunidade(dias, status, valor, cl["categoria"], eh_conc, esfera, uf, qual_frac)
    prio = faixa_prioridade(score)
    # oportunidade derivada
    opp = ("Sim" if (cl["eh_ti"] == "Sim" and valor >= CONFIG["OPP_VALOR_MIN"]
                     and (status == "Vencido" or (dias is not None and dias <= CONFIG["OPP_DIAS_MAX"])))
           else "Não")
    # link público do contrato
    cnpj_org = limpa_cnpj(org.get("cnpj"))
    ano = it.get("anoContrato"); seq = it.get("sequencialContrato")
    link = f"https://pncp.gov.br/app/contratos/{cnpj_org}/{ano}/{seq}" if (cnpj_org and ano and seq) else "https://pncp.gov.br"
    nome_org = org.get("razaoSocial") or uni.get("nomeUnidade") or ""

    return {
        # Identificação
        "ID Interno": f"PNCP-{idx:05d}",
        "Fonte": "PNCP",
        "Data da Coleta": HOJE.strftime("%d/%m/%Y"),
        "Link da Fonte": link,
        "Número do Contrato": it.get("numeroContratoEmpenho") or "",
        "Número do Processo": it.get("processo") or "",
        "ID PNCP": it.get("numeroControlePNCP") or it.get("numeroControlePncpCompra") or "",
        "Modalidade": (it.get("tipoContrato") or {}).get("nome") or "",
        "Situação": (it.get("categoriaProcesso") or {}).get("nome") or "",
        # Órgão
        "Órgão": nome_org,
        "Nome Padronizado (Órgão)": pad_nome(nome_org),
        "CNPJ do Órgão": formata_cnpj(org.get("cnpj")),
        "UF": uf,
        "Município": uni.get("municipioNome") or "",
        "Esfera": esfera,
        "Poder": poder,
        "Tipo de Órgão": "",  # PNCP não traz; inferível depois
        "Segmento Presumido": segmento_presumido(nome_org),
        # Fornecedor
        "Fornecedor": forn_nome,
        "Nome Padronizado (Fornecedor)": pad_nome(forn_nome),
        "CNPJ do Fornecedor": formata_cnpj(it.get("niFornecedor")),
        "Tipo Fornecedor": it.get("tipoPessoa") or "",
        "Possível Concorrente": "Sim" if eh_conc else "Não",
        "Grau de Ameaça": grau,
        # Contrato
        "Objeto": objeto.strip(),
        "Objeto Normalizado": norm(objeto)[:300],
        "Categoria Principal": cl["categoria"],
        "Subcategoria": cl["subcategoria"],
        "Palavras-chave Encontradas": ", ".join(cl["palavras"][:8]),
        "Valor Total": round(valor, 2),
        "Valor Mensal Estimado": round(valor / max(1, meses(inicio, fim)), 2) if (valor and fim and inicio) else "",
        "Data de Assinatura": fmt_data(assinatura),
        "Data de Início": fmt_data(inicio),
        "Data de Término": fmt_data(fim),
        "Dias até Vencimento": dias if dias is not None else "",
        "Status do Contrato": status,
        "Possibilidade de Renovação": "Alta" if status in ("Vencido","Vencendo em 30 dias","Vencendo em 60 dias") else ("Média" if status == "Vencendo em 90 dias" else "Baixa"),
        "Link do Contrato": link,
        "Link do Edital": "",
        # Inteligência comercial
        "É TI?": cl["eh_ti"],
        "Confiança da Classificação": cl["confianca"],
        "Oportunidade Derivada": opp,
        "Score de Oportunidade": score,
        "Prioridade": prio,
        "Motivo da Prioridade": motivo_prioridade(score, dias, status, valor, cl["categoria"], eh_conc),
        "Próxima Ação Recomendada": proxima_acao(cl["eh_ti"], status, dias),
        "Responsável Comercial Sugerido": CONFIG["RESP_POR_UF"].get(uf, "A definir"),
        "Carteira Sugerida": CONFIG["CARTEIRA_POR_UF"].get(uf, uf),
        "Qualidade do Dado": qualidade,
        "Observações": "",
        "Necessita Revisão?": cl["revisao"],
        # chave de dedup (interna)
        "_chave": it.get("numeroControlePNCP") or f"{cnpj_org}-{ano}-{seq}-{it.get('numeroContratoEmpenho')}",
    }

def meses(ini, fim):
    if not ini or not fim: return 1
    return max(1, round((fim - ini).days / 30))

def segmento_presumido(nome_org):
    n = norm(nome_org)
    mapa = [("saude","Saúde"),("educac","Educação"),("seguranca","Segurança Pública"),
            ("fazenda","Fazenda/Tributação"),("tribunal","Justiça"),("justic","Justiça"),
            ("transito","Transporte"),("transporte","Transporte"),("ambiente","Meio Ambiente"),
            ("agricultura","Agricultura"),("camara","Legislativo"),("assembleia","Legislativo")]
    for k, v in mapa:
        if _tem(k, n) or k in n: return v
    return "Administração"

def processar(registros):
    log("Processando, classificando e pontuando...")
    linhas, vistos = [], set()
    dups = 0
    for i, it in enumerate(registros, 1):
        linha = construir_linha(it, i)
        ch = linha["_chave"]
        if ch in vistos:
            dups += 1; continue
        vistos.add(ch)
        if linha["Valor Total"] < CONFIG["VALOR_MIN"]:
            continue
        linhas.append(linha)
    log(f"   linhas únicas: {len(linhas)} | duplicatas removidas: {dups}")
    return linhas

# =============================== SAÍDAS ===============================
def _csv(path, linhas, cols):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for l in linhas: w.writerow(l)
    log(f"   -> {os.path.basename(path)} ({len(linhas)} linhas)")

def salvar_saidas(linhas):
    od = CONFIG["OUT_DIR"]; os.makedirs(od, exist_ok=True)
    cols = [c for c in linhas[0].keys() if c != "_chave"] if linhas else []
    # bruto (todos DF/GO, antes do filtro de TI)
    _csv(os.path.join(od, "atlas_pncp_contratos_bruto.csv"), linhas, cols)
    # tratado: só TI
    ti = [l for l in linhas if l["É TI?"] == "Sim"]
    _csv(os.path.join(od, "atlas_pncp_contratos_ti_tratado.csv"), ti, cols)
    # oportunidades
    opp = sorted([l for l in ti if l["Oportunidade Derivada"] == "Sim"],
                 key=lambda x: x["Score de Oportunidade"], reverse=True)
    _csv(os.path.join(od, "atlas_pncp_oportunidades.csv"), opp, cols)
    return ti, opp

# =============================== RELATÓRIO ===============================
def relatorio(linhas, ti, opp, total_nac):
    od = CONFIG["OUT_DIR"]
    venc = lambda lo, hi: [l for l in ti if isinstance(l["Dias até Vencimento"], int) and lo <= l["Dias até Vencimento"] <= hi]
    v30, v60, v90 = venc(0,30), venc(31,60), venc(61,90)
    vencidos = [l for l in ti if l["Status do Contrato"] == "Vencido"]
    valor_total = sum(l["Valor Total"] for l in ti)
    por_uf  = collections.Counter(l["UF"] for l in ti)
    por_cat = collections.Counter(l["Categoria Principal"] for l in ti)
    por_forn= collections.Counter(l["Fornecedor"] for l in ti if l["Fornecedor"])
    por_org = collections.Counter(l["Órgão"] for l in ti if l["Órgão"])
    baixa_q = [l for l in ti if l["Qualidade do Dado"] == "Baixa" or l["Necessita Revisão?"] == "Sim"]
    criticas= [l for l in opp if l["Prioridade"] in ("Máxima","Alta")]
    top_valor = sorted(ti, key=lambda x: x["Valor Total"], reverse=True)[:10]
    top_venc  = sorted([l for l in ti if isinstance(l["Dias até Vencimento"], int) and l["Dias até Vencimento"]>=0],
                       key=lambda x: x["Dias até Vencimento"])[:10]
    concs = collections.Counter(l["Fornecedor"] for l in ti if l["Possível Concorrente"]=="Sim" and l["Fornecedor"])

    # ---- resumo textual (também vai pro log)
    log("="*64); log("RELATÓRIO EXECUTIVO DA INGESTÃO")
    log(f"Janela de publicação: {CONFIG['DATA_INICIAL']}–{CONFIG['DATA_FINAL']} | UFs: {CONFIG['UFS']}")
    log(f"Registros DF/GO coletados: {len(linhas)} (de ~{total_nac} nacionais na janela)")
    log(f"Contratos classificados como TI: {len(ti)} | descartados (não-TI): {len(linhas)-len(ti)}")
    log(f"Oportunidades derivadas: {len(opp)} | críticas (Máxima/Alta): {len(criticas)}")
    log(f"Valor total mapeado (TI): {brl(valor_total)}")
    log(f"Vencendo 30d: {len(v30)} | 60d: {len(v60)} | 90d: {len(v90)} | vencidos: {len(vencidos)}")
    log(f"Por UF: {dict(por_uf)}")
    log(f"Top categorias: {por_cat.most_common(6)}")
    log(f"Top fornecedores: {por_forn.most_common(5)}")
    log(f"Possíveis concorrentes: {concs.most_common(5)}")
    log(f"Baixa qualidade / revisão: {len(baixa_q)}")
    if opp:
        log("TOP 5 OPORTUNIDADES:")
        for l in opp[:5]:
            log(f"   [{l['Score de Oportunidade']}|{l['Prioridade']}] {l['Nome Padronizado (Órgão)']} ({l['UF']}) "
                f"· {l['Categoria Principal']} · {brl(l['Valor Total'])} · {l['Status do Contrato']} · {l['Próxima Ação Recomendada']}")
    log("="*64)

    # ---- XLSX
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "Resumo"
        H = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1E293B")
        def head(ws, row, cols):
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row, i, c); cell.font = H; cell.fill = FILL; cell.alignment = Alignment(horizontal="left")
        kv = [
            ("Janela de publicação", f"{CONFIG['DATA_INICIAL']}–{CONFIG['DATA_FINAL']}"),
            ("UFs", ", ".join(CONFIG["UFS"])),
            ("Registros DF/GO coletados", len(linhas)),
            ("Total nacional na janela (~)", total_nac),
            ("Contratos de TI", len(ti)),
            ("Descartados (não-TI)", len(linhas)-len(ti)),
            ("Oportunidades derivadas", len(opp)),
            ("Oportunidades críticas (Máxima/Alta)", len(criticas)),
            ("Valor total mapeado (TI)", valor_total),
            ("Vencendo em 30 dias", len(v30)),
            ("Vencendo em 60 dias", len(v60)),
            ("Vencendo em 90 dias", len(v90)),
            ("Contratos vencidos", len(vencidos)),
            ("Baixa qualidade / revisão", len(baixa_q)),
            ("Data da coleta", HOJE.strftime("%d/%m/%Y")),
        ]
        head(ws, 1, ["Indicador", "Valor"])
        for r, (k, v) in enumerate(kv, 2):
            ws.cell(r, 1, k); ws.cell(r, 2, v)
        ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 26

        def aba_contagem(nome, counter, c1):
            w = wb.create_sheet(nome); head(w, 1, [c1, "Qtd"])
            for r, (k, v) in enumerate(counter.most_common(50), 2):
                w.cell(r, 1, k or "—"); w.cell(r, 2, v)
            w.column_dimensions["A"].width = 46; w.column_dimensions["B"].width = 10
        aba_contagem("Por UF", por_uf, "UF")
        aba_contagem("Por Categoria", por_cat, "Categoria")
        aba_contagem("Por Fornecedor", por_forn, "Fornecedor")
        aba_contagem("Por Órgão", por_org, "Órgão")
        aba_contagem("Possíveis Concorrentes", concs, "Fornecedor")

        def aba_lista(nome, linhas_, campos):
            w = wb.create_sheet(nome); head(w, 1, campos)
            for r, l in enumerate(linhas_, 2):
                for i, c in enumerate(campos, 1): w.cell(r, i, l.get(c, ""))
            for i, c in enumerate(campos, 1):
                w.column_dimensions[chr(64+i)].width = 16 if c not in ("Órgão","Objeto","Motivo da Prioridade","Próxima Ação Recomendada") else 42
        campos_opp = ["Score de Oportunidade","Prioridade","Nome Padronizado (Órgão)","UF","Categoria Principal",
                      "Valor Total","Status do Contrato","Dias até Vencimento","Fornecedor","Próxima Ação Recomendada"]
        aba_lista("Top Oportunidades", opp[:30], campos_opp)
        aba_lista("Top por Valor", top_valor, campos_opp)
        aba_lista("Vencendo primeiro", top_venc, campos_opp)
        aba_lista("Baixa Qualidade", baixa_q[:50], ["ID Interno","Órgão","UF","Objeto","Necessita Revisão?","Qualidade do Dado"])
        wb.save(os.path.join(od, "atlas_pncp_relatorio_ingestao.xlsx"))
        log("   -> atlas_pncp_relatorio_ingestao.xlsx")
    except Exception as e:
        log(f"   (relatório xlsx não gerado: {e})")

    # ---- log txt
    with open(os.path.join(od, "atlas_pncp_log_execucao.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(_LOG))
    log("   -> atlas_pncp_log_execucao.txt")

# =============================== MAIN ===============================
def main():
    t0 = time.time()
    log("ATLAS B2G — Ingestão PNCP (piloto DF+GO+TI) iniciada.")
    registros, total_nac = coletar()
    if not registros:
        log("Nenhum registro DF/GO na janela. Ajuste DATA_INICIAL/DATA_FINAL/MAX_PAGINAS no CONFIG.")
        # ainda assim grava log
        os.makedirs(CONFIG["OUT_DIR"], exist_ok=True)
        with open(os.path.join(CONFIG["OUT_DIR"], "atlas_pncp_log_execucao.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(_LOG))
        return
    linhas = processar(registros)
    if not linhas:
        log("Sem linhas após filtros."); return
    ti, opp = salvar_saidas(linhas)
    relatorio(linhas, ti, opp, total_nac)
    log(f"Concluído em {time.time()-t0:.1f}s.")

if __name__ == "__main__":
    main()
