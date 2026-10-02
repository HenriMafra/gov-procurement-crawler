# -*- coding: utf-8 -*-
"""
ATLAS B2G — Orquestrador da Rotina Semanal
==========================================
Executa a esteira completa de inteligência comercial B2G:
config -> ingestão PNCP -> classificação/score -> lista de ataque -> comparativo
com a rodada anterior -> Excel executivo -> relatório -> protótipo -> pacote ZIP.

Uso:
  python src/atlas_weekly_runner.py --config config/atlas_config.json
  python src/atlas_weekly_runner.py --ufs DF GO --date 2026-05-31 --no-cache --score-min 70
"""
import os, sys, json, csv, time, datetime, argparse, shutil, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
import utils
import atlas_pncp_ingest as core
import atlas_ataque_comercial as atk

RUN_LOG = []
def log(m):
    line = f"[{datetime.datetime.now():%Y-%m-%d %H:%M:%S}] {m}"; RUN_LOG.append(line); print(line, flush=True)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))   # ATLAS-PNCP-Pilot/
def P(*a): return os.path.join(ROOT, *a)

# ---------------- CLI + CONFIG ----------------
def parse_args():
    ap = argparse.ArgumentParser(description="ATLAS B2G — rotina semanal")
    ap.add_argument("--config", default=P("config", "atlas_config.json"))
    ap.add_argument("--date", help="Data da rodada (YYYY-MM-DD). Padrão: hoje.")
    ap.add_argument("--ufs", nargs="*", help="UFs (ex.: DF GO)")
    ap.add_argument("--no-cache", action="store_true")
    ap.add_argument("--score-min", type=int)
    ap.add_argument("--max-pag", type=int)
    ap.add_argument("--tag", help="Rótulo da rodada (ex.: PRODUCAO) — prefixa a pasta e nomeia o ZIP.")
    ap.add_argument("--write-db", action="store_true", help="Grava a rodada no banco (além de gerar os arquivos).")
    ap.add_argument("--db-url", help="URL do banco (senão usa $DATABASE_URL ou sqlite local de teste).")
    ap.add_argument("--no-remove", action="store_true", help="ADD-only: nunca marca oportunidades como 'Removida' (a base só cresce). Usar na coleta pinada.")
    ap.add_argument("--no-init", action="store_true", help="Não re-inicializa schema/seed na gravação (evita corrida de DDL no pool concorrente e preserva a taxonomia).")
    return ap.parse_args()

def aplicar_config(cfg, args):
    if args.ufs: cfg["ufs"] = args.ufs
    if args.no_cache: cfg["usar_cache"] = False
    if args.score_min is not None: cfg["score_minimo"] = args.score_min
    if args.max_pag is not None: cfg["max_paginas_por_janela"] = args.max_pag
    return cfg

def montar_janelas(cfg, run_date):
    j = cfg["janelas"]
    modo = j.get("modo")
    # ULTIMOS_ANOS: janela deslizante de N anos-calendário (ano atual + anteriores),
    # avança sozinha pra sempre. Ex.: em 2026 => 2026..2021; em 2039 => 2039..2034.
    # Ano atual vai de 01/01 até hoje; anos anteriores 01/01..31/12 (respeita limite
    # de 365 dias por consulta do PNCP). Mais recente primeiro.
    if modo == "ultimos_anos":
        n = int(j.get("anos", 6))
        Y = run_date.year
        out = []
        for y in range(Y, Y - n, -1):
            ini = f"{y}0101"
            fim = run_date.strftime("%Y%m%d") if y == Y else f"{y}1231"
            out.append((ini, fim))
        return out
    if modo == "relativo":
        out = []
        for ini, fim in j["relativo_dias"]:
            out.append([(run_date - datetime.timedelta(days=ini)).strftime("%Y%m%d"),
                        (run_date - datetime.timedelta(days=fim)).strftime("%Y%m%d")])
        return [tuple(x) for x in out]
    return [tuple(x) for x in j["explicito"]]

# ---------------- COMPARATIVO ----------------
def comparar(L, rodadas_dir, data_atual, tag=None):
    """Anota cada oportunidade com status_na_rodada e mudanças vs rodada anterior (mesmo rótulo)."""
    import re
    prev_map, prev_date = {}, None
    if os.path.isdir(rodadas_dir):
        pref = (tag + "_") if tag else ""
        pat = re.compile(r"^" + re.escape(pref) + r"(\d{4}-\d{2}-\d{2})$")
        datas = [m.group(1) for m in (pat.match(n) for n in os.listdir(rodadas_dir)) if m and m.group(1) < data_atual]
        for d in sorted(datas, reverse=True):
            cand = os.path.join(rodadas_dir, pref + d, f"atlas_lista_ataque_comercial_{d}.csv")
            rows = utils.ler_csv(cand)
            if rows:
                prev_map = {r.get("ID PNCP"): r for r in rows if r.get("ID PNCP")}
                prev_date = d; break
    ids_atual = set(l.get("ID PNCP") for l in L)
    for l in L:
        if not prev_map:
            l["Status na Rodada"] = "Primeira rodada"; l["Mudança Score"] = ""; l["Mudança Urgência"] = ""
            l["Mudança Valor"] = ""; l["Observação Comparativo"] = "Sem rodada anterior para comparar."
            continue
        p = prev_map.get(l.get("ID PNCP"))
        if not p:
            l["Status na Rodada"] = "Nova"; l["Mudança Score"] = ""; l["Mudança Urgência"] = ""
            l["Mudança Valor"] = ""; l["Observação Comparativo"] = "Nova oportunidade nesta rodada."
        else:
            ds = l["Score de Oportunidade"] - int(float(p.get("Score de Oportunidade") or 0))
            du = (p.get("Urgência Comercial") or "") != l["Urgência Comercial"]
            dv = l["Valor Total"] - float(p.get("Valor Total") or 0)
            dst = (p.get("Status do Contrato") or "") != l["Status do Contrato"]
            mudou = ds != 0 or du or dst or abs(dv) > 1
            l["Status na Rodada"] = "Alterada" if mudou else "Mantida"
            l["Mudança Score"] = f"{'+' if ds>=0 else ''}{ds}" if ds else "0"
            l["Mudança Urgência"] = f"{p.get('Urgência Comercial')}→{l['Urgência Comercial']}" if du else ""
            l["Mudança Valor"] = round(dv, 2) if abs(dv) > 1 else 0
            obs = []
            if du: obs.append("urgência mudou")
            if dst: obs.append("status mudou")
            if ds: obs.append("score " + l["Mudança Score"])
            l["Observação Comparativo"] = "; ".join(obs) if obs else "mantida"
    removidas = [p for k, p in prev_map.items() if k not in ids_atual] if prev_map else []
    return prev_date, removidas

# ---------------- PROTÓTIPO ----------------
def to_js(l):
    g = lambda k: l.get(k, "")
    def num(k):
        try: return float(l.get(k) or 0)
        except Exception: return 0
    dias = l.get("Dias até Vencimento")
    try: dias = int(dias)
    except Exception: dias = None
    return {
        "id": g("ID Oportunidade"), "pncp": g("ID PNCP"), "link": g("Link da Fonte") or "https://pncp.gov.br",
        "numContrato": g("Número do Contrato"), "processo": g("Número do Processo"),
        "orgao": g("Órgão"), "orgaoPad": g("Nome Padronizado (Órgão)") or g("Órgão"), "cnpjOrgao": g("CNPJ do Órgão"),
        "uf": g("UF"), "municipio": g("Município"), "poder": g("Poder"), "esfera": g("Esfera"),
        "unidade": g("Unidade Compradora"), "segmento": g("Segmento Presumido"),
        "objeto": g("Objeto"), "subcat": g("Subcategoria"), "palavras": g("Palavras-chave Encontradas"),
        "categoria": g("Categoria Principal"), "valor": num("Valor Total"), "valorMensal": g("Valor Mensal Estimado"),
        "assinatura": g("Data de Assinatura"), "inicio": g("Início da Vigência"), "fim": g("Fim da Vigência"),
        "dias": dias, "status": g("Status do Contrato"), "janela": g("Janela Comercial"),
        "fornecedor": g("Fornecedor"), "fornecedorPad": g("Nome Padronizado (Fornecedor)"), "cnpjForn": g("CNPJ do Fornecedor"),
        "concorrente": g("Possível Concorrente") == "Sim", "ameaca": g("Grau de Ameaça"),
        "recorrencia": int(float(l.get("Recorrência do Fornecedor") or 1)),
        "conf": g("Confiança da Classificação"), "score": int(l.get("Score de Oportunidade") or 0),
        "prioridade": g("Prioridade"), "motivo": g("Motivo da Prioridade"), "tipo": g("Tipo de Oportunidade"),
        "proximaAcao": g("Próxima Ação Recomendada"), "argumento": g("Argumento Comercial Sugerido"),
        "urgencia": g("Urgência Comercial"), "responsavel": g("Responsável Comercial Sugerido"), "carteira": g("Carteira Sugerida"),
        "revisao": g("Necessita Revisão?") == "Sim", "obs": g("Observações"), "explica": g("explicacao_score") or g("Motivo da Prioridade"),
    }

def gerar_prototipo(L, html_path, js_path):
    data = [to_js(l) for l in L]
    js_arr = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    open(js_path, "w", encoding="utf-8").write("window.ATLAS_DATA = " + js_arr + ";\n")
    tpl = open(P("templates", "prototipo_template.html"), encoding="utf-8").read()
    open(html_path, "w", encoding="utf-8").write(tpl.replace("/*__ATLAS_DATA__*/", "window.ATLAS_DATA = " + js_arr + ";"))

# ---------------- EXCEL SEMANAL ----------------
def gerar_excel(path, L, agg, removidas, prev_date, params):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); HF = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="0F172A")
    thin = Side(style="thin", color="E2E8F0"); BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
    def head(ws, row, cols, w=None):
        for i, c in enumerate(cols, 1):
            x = ws.cell(row, i, c); x.font = HF; x.fill = HFILL; x.alignment = Alignment(vertical="center")
            if w: ws.column_dimensions[get_column_letter(i)].width = w[i-1]
        ws.row_dimensions[row].height = 20
    novas = [l for l in L if l.get("Status na Rodada") == "Nova"]
    alteradas = [l for l in L if l.get("Status na Rodada") == "Alterada"]
    # 1 Resumo
    ws = wb.active; ws.title = "Resumo Executivo"; ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:E1"); ws["A1"] = "ATLAS B2G — Lista de Ataque Comercial (rodada semanal)"; ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"DF+GO · Tecnologia · {params['data']} · fonte PNCP (dados públicos)"; ws["A2"].font = Font(italic=True, color="64748B")
    kp = [("Oportunidades", agg["n"]), ("Críticas", agg["criticas"]), ("Valor total", utils.brl(agg["valor"])),
          ("Ataque imediato", agg.get("imediato", "")), ("Vencidos recentes", agg["vencidos"]), ("Vencendo ≤30d", agg["v30"]),
          ("DF", agg["uf"].get("DF", 0)), ("GO", agg["uf"].get("GO", 0)),
          ("Novas vs rodada anterior", len(novas) if prev_date else "—"),
          ("Alteradas", len(alteradas) if prev_date else "—"), ("Removidas", len(removidas) if prev_date else "—"),
          ("Rodada anterior", prev_date or "primeira rodada")]
    head(ws, 4, ["Indicador", "Valor"], [34, 26])
    for r, (k, v) in enumerate(kp, 5): ws.cell(r, 1, k).border = BORD; ws.cell(r, 2, v).border = BORD
    for cc, t in ((4, "Top 10 — Órgão"), (5, "Score")): x = ws.cell(4, cc, t); x.font = HF; x.fill = HFILL
    ws.column_dimensions["D"].width = 46; ws.column_dimensions["E"].width = 8
    for r, l in enumerate(L[:10], 5):
        ws.cell(r, 4, f"{l['Nome Padronizado (Órgão)']} ({l['UF']}) · {l['Categoria Principal']}").border = BORD
        ws.cell(r, 5, l["Score de Oportunidade"]).border = BORD
    # 2 Lista de Ataque
    cols = [c for c in (L[0].keys() if L else []) if c != "explicacao_score"]
    wl = wb.create_sheet("Lista de Ataque"); head(wl, 1, cols)
    for r, l in enumerate(L, 2):
        for i, c in enumerate(cols, 1):
            x = wl.cell(r, i, l.get(c, "")); x.border = BORD
            if c == "Prioridade": x.fill = PatternFill("solid", fgColor=atk.PRIO_FILL.get(l[c], "FFFFFF")); x.font = Font(bold=True, color="FFFFFF")
            if c == "Urgência Comercial": x.fill = PatternFill("solid", fgColor=atk.URG_FILL.get(l[c], "FFFFFF")); x.font = Font(bold=True, color="FFFFFF")
            if c == "Valor Total": x.number_format = '"R$" #,##0'
    wide = {"Órgão": 32, "Nome Padronizado (Órgão)": 30, "Objeto": 48, "Motivo da Prioridade": 44, "Próxima Ação Recomendada": 44, "Argumento Comercial Sugerido": 46, "Fornecedor": 26}
    for i, c in enumerate(cols, 1):
        ws_w = wide.get(c, 15); wl.column_dimensions[get_column_letter(i)].width = ws_w
    wl.freeze_panes = "A2"; wl.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(L)+1}"
    # 3 Top 10
    tc = ["Score de Oportunidade", "Prioridade", "Urgência Comercial", "Nome Padronizado (Órgão)", "UF", "Categoria Principal", "Valor Total", "Status do Contrato", "Próxima Ação Recomendada"]
    wt = wb.create_sheet("Top 10 Semana"); head(wt, 1, tc, [8, 16, 12, 30, 6, 16, 14, 16, 46])
    for r, l in enumerate(L[:10], 2):
        for i, c in enumerate(tc, 1):
            x = wt.cell(r, i, l.get(c, "")); x.border = BORD
            if c == "Prioridade": x.fill = PatternFill("solid", fgColor=atk.PRIO_FILL.get(l[c], "FFFFFF")); x.font = Font(bold=True, color="FFFFFF")
            if c == "Valor Total": x.number_format = '"R$" #,##0'
    # 4 Por Vendedor
    porv = collections.defaultdict(list)
    for l in L: porv[l["Responsável Comercial Sugerido"]].append(l)
    wv = wb.create_sheet("Por Responsável"); head(wv, 1, ["Responsável sugerido", "Oport.", "Críticas", "Valor total"], [34, 10, 10, 18])
    for r, (k, ops) in enumerate(sorted(porv.items(), key=lambda x: -sum(o["Valor Total"] for o in x[1])), 2):
        wv.cell(r, 1, k).border = BORD; wv.cell(r, 2, len(ops)).border = BORD
        wv.cell(r, 3, sum(1 for o in ops if o["Urgência Comercial"] == "Crítica")).border = BORD
        c = wv.cell(r, 4, round(sum(o["Valor Total"] for o in ops))); c.number_format = '"R$" #,##0'; c.border = BORD
    # 5/6/7 contagens
    def aba(nome, counter, c1, vmap=None):
        w = wb.create_sheet(nome)
        head(w, 1, [c1, "Qtd", "Valor"] if vmap else [c1, "Qtd"], [44, 10, 18] if vmap else [44, 10])
        for r, (k, v) in enumerate(counter.most_common(60), 2):
            w.cell(r, 1, k or "—").border = BORD; w.cell(r, 2, v).border = BORD
            if vmap: c = w.cell(r, 3, round(vmap.get(k, 0))); c.number_format = '"R$" #,##0'; c.border = BORD
    aba("Por UF", agg["uf"], "UF", agg["uf_valor"]); aba("Por Categoria", agg["cat"], "Categoria", agg["cat_valor"]); aba("Por Fornecedor", agg["forn"], "Fornecedor", agg["forn_valor"])
    conc_c = collections.Counter(l["Fornecedor"] for l in L if l["Possível Concorrente"] == "Sim" and l["Fornecedor"])
    conc_v = collections.defaultdict(float)
    for l in L:
        if l["Possível Concorrente"] == "Sim" and l["Fornecedor"]: conc_v[l["Fornecedor"]] += l["Valor Total"]
    aba("Por Concorrente", conc_c, "Fornecedor (possível concorrente)", conc_v)
    # 8 Revisão
    wr = wb.create_sheet("Revisão Manual"); rc = ["ID Oportunidade", "Órgão", "UF", "Valor Total", "Confiança da Classificação", "Necessita Revisão?", "Fim da Vigência", "Objeto"]
    head(wr, 1, rc, [14, 32, 6, 14, 14, 14, 14, 56])
    rev = [l for l in L if l["Necessita Revisão?"] == "Sim" or l["Confiança da Classificação"] == "Baixa" or l["Fim da Vigência"] == ""]
    for r, l in enumerate(rev, 2):
        for i, c in enumerate(rc, 1):
            x = wr.cell(r, i, l.get(c, "")); x.border = BORD
            if c == "Valor Total": x.number_format = '"R$" #,##0'
    # 9 Comparativo Semanal
    wc = wb.create_sheet("Comparativo Semanal"); cc = ["Status na Rodada", "Nome Padronizado (Órgão)", "UF", "Categoria Principal", "Valor Total", "Mudança Score", "Mudança Urgência", "Observação Comparativo"]
    head(wc, 1, cc, [16, 30, 6, 16, 14, 12, 22, 30])
    comp_rows = [l for l in L if l.get("Status na Rodada") in ("Nova", "Alterada", "Primeira rodada")]
    for r, l in enumerate(comp_rows, 2):
        for i, c in enumerate(cc, 1):
            x = wc.cell(r, i, l.get(c, "")); x.border = BORD
            if c == "Valor Total": x.number_format = '"R$" #,##0'
    base = len(comp_rows) + 3
    wc.cell(base, 1, "REMOVIDAS (saíram vs rodada anterior)").font = Font(bold=True, color="DC2626")
    for r, p in enumerate(removidas, base + 1):
        wc.cell(r, 1, "Removida"); wc.cell(r, 2, p.get("Nome Padronizado (Órgão)", "")); wc.cell(r, 3, p.get("UF", "")); wc.cell(r, 4, p.get("Categoria Principal", ""))
    # 10 Parâmetros
    wp = wb.create_sheet("Parâmetros"); head(wp, 1, ["Parâmetro", "Valor"], [34, 80])
    for r, (k, v) in enumerate(params.items(), 2): wp.cell(r, 1, k).border = BORD; wp.cell(r, 2, str(v)).border = BORD
    wb.save(path)

# ---------------- RELATÓRIO ----------------
def gerar_relatorio(path_md, L, agg, removidas, prev_date, params, segundos):
    novas = [l for l in L if l.get("Status na Rodada") == "Nova"]
    alteradas = [l for l in L if l.get("Status na Rodada") == "Alterada"]
    rev = [l for l in L if l["Necessita Revisão?"] == "Sim" or l["Confiança da Classificação"] == "Baixa" or l["Fim da Vigência"] == ""]
    def linha(l):
        v = atk.venc_legivel(l["Dias até Vencimento"] if isinstance(l["Dias até Vencimento"], int) else None, atk._parse_fim(l["Fim da Vigência"]))
        return f"- **{l['Nome Padronizado (Órgão)']}** ({l['UF']}) · {l['Categoria Principal']} · {utils.brl(l['Valor Total'])} · {v} · _{l['Prioridade']}_ — {l['Próxima Ação Recomendada']}"
    m = []
    m.append(f"# ATLAS B2G — Relatório Semanal de Inteligência Comercial\n")
    m.append(f"**Rodada {params['data']} · DF + Goiás · Tecnologia · fonte PNCP (dados públicos)**\n")
    m.append("## 1. Resumo executivo")
    m.append(f"- Oportunidades de TI na janela comercial: **{agg['n']}** · críticas: **{agg['criticas']}** · valor total: **{utils.brl(agg['valor'])}**")
    m.append(f"- Ataque imediato: **{agg.get('imediato',0)}** · vencidos recentes: **{agg['vencidos']}** · vencendo ≤30d: **{agg['v30']}**")
    m.append(f"- Distribuição: DF **{agg['uf'].get('DF',0)}** · GO **{agg['uf'].get('GO',0)}**")
    if prev_date: m.append(f"- Comparado à rodada {prev_date}: **{len(novas)} novas**, **{len(alteradas)} alteradas**, **{len(removidas)} removidas**.")
    else: m.append("- Primeira rodada (sem comparativo).")
    m.append(f"- Tempo de execução: **{segundos:.0f}s**.\n")
    m.append("**Recomendação da semana:** priorizar as oportunidades *Crítica/Ataque imediato* abaixo e distribuir o Top 10 para os responsáveis sugeridos.\n")
    m.append("## 2. Top 10 oportunidades")
    for l in L[:10]: m.append(linha(l))
    m.append("\n## 3. Mudanças vs rodada anterior")
    if prev_date:
        m.append(f"- **Novas ({len(novas)}):** " + (", ".join(f"{l['Nome Padronizado (Órgão)']} ({l['UF']})" for l in novas[:10]) or "—"))
        m.append(f"- **Alteradas ({len(alteradas)}):** " + (", ".join(f"{l['Nome Padronizado (Órgão)']} [{l['Mudança Urgência'] or l['Mudança Score']}]" for l in alteradas[:10]) or "—"))
        m.append(f"- **Removidas ({len(removidas)}):** " + (", ".join(p.get('Nome Padronizado (Órgão)', '') for p in removidas[:10]) or "—"))
    else:
        m.append("- Sem rodada anterior.")
    m.append("\n## 4. Visão por UF")
    for uf in params["ufs"]:
        m.append(f"- **{uf}**: {agg['uf'].get(uf,0)} oport. · {utils.brl(agg['uf_valor'].get(uf,0))} · top: " + ", ".join(f"{k}({v})" for k, v in agg["uf_cat"].get(uf, collections.Counter()).most_common(3)))
    m.append("\n## 5. Visão por categoria")
    for k, v in agg["cat"].most_common(12): m.append(f"- {k}: {v} · {utils.brl(agg['cat_valor'].get(k,0))}")
    m.append("\n## 6. Visão por fornecedor")
    for k, v in agg["forn"].most_common(8):
        flag = " — possível concorrente" if k in agg["concorrentes"] else ""
        m.append(f"- {k}: {v} contrato(s) · {utils.brl(agg['forn_valor'].get(k,0))}{flag}")
    m.append("\n## 7. Lista de ataque da semana (priorizar)")
    m.append("| # | Órgão | UF | Categoria | Valor | Situação | Urgência | Ação |")
    m.append("|---|---|---|---|---|---|---|---|")
    for i, l in enumerate(L[:10], 1):
        v = atk.venc_legivel(l["Dias até Vencimento"] if isinstance(l["Dias até Vencimento"], int) else None, atk._parse_fim(l["Fim da Vigência"]))
        m.append(f"| {i} | {l['Nome Padronizado (Órgão)']} | {l['UF']} | {l['Categoria Principal']} | {utils.brl(l['Valor Total'])} | {v} | {l['Urgência Comercial']} | {l['Próxima Ação Recomendada'][:54]} |")
    m.append("\n## 8. Pendências da operação")
    m.append(f"- Itens que pedem validação humana: **{len(rev)}** (objeto ambíguo / baixa confiança / sem data).")
    m.append(f"- Sem responsável definido: **{sum(1 for l in L if 'A definir' in l['Responsável Comercial Sugerido'])}** (todas — atribuir na coordenação).")
    m.append("\n## 9. Recomendações")
    m.append("- Distribuir o **Top 10** para os responsáveis sugeridos por UF.")
    m.append("- Validar primeiro os itens da aba **Revisão Manual** do Excel.")
    m.append("- Cadastrar concorrentes reais em `config/atlas_config.json` → `concorrentes_conhecidos`.")
    m.append("- Priorizar categorias estratégicas (Cyber, Infra, Cloud, Dados/IA).")
    m.append("\n---\n*Dados públicos do PNCP. Fornecedores marcados como 'possível concorrente' a validar. Itens 'necessita revisão' dependem de validação humana. Nenhuma oportunidade é garantida.*")
    open(path_md, "w", encoding="utf-8").write("\n".join(m))
    # PDF opcional
    pdf = path_md.replace(".md", ".pdf")
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import cm
        c = canvas.Canvas(pdf, pagesize=A4); w, h = A4; y = h - 2 * cm
        for ln in "\n".join(m).split("\n"):
            ln = ln.replace("**", "").replace("`", "").replace("#", "").strip()
            if y < 2 * cm: c.showPage(); y = h - 2 * cm
            c.setFont("Helvetica", 9); c.drawString(2 * cm, y, ln[:110]); y -= 0.42 * cm
        c.save(); return pdf, True
    except Exception:
        return None, False

# ---------------- PRODUÇÃO (calibração comercial) ----------------
def grupo_categoria(cat):
    if cat == "Cibersegurança": return "cyber"
    if cat in ("Software", "Dados/BI/IA"): return "software"
    if cat in ("Serviços/Outsourcing", "Governança de TI"): return "servicos"
    if cat in ("Infraestrutura", "Cloud", "Redes", "Hardware", "Backup/DR", "CFTV/Videomonitoramento", "Telecom"): return "infra"
    return "default"

def atribuir_responsavel(l, cfg):
    uf = l.get("UF", ""); g = grupo_categoria(l.get("Categoria Principal", ""))
    rmap = (cfg.get("responsaveis") or {}).get(uf) or {}
    return rmap.get(g) or rmap.get("default") or (cfg.get("responsavel_por_uf") or {}).get(uf) or "A definir"

def enriquecer_producao(L, cfg):
    """Ajustes de produção: bônus de recorrência no score, responsável por UF×categoria, motivo da ação."""
    for l in L:
        try: rec = int(float(l.get("Recorrência do Fornecedor") or 1))
        except Exception: rec = 1
        bonus = 4 if rec >= 3 else (2 if rec >= 2 else 0)
        sc = min(100, int(l["Score de Oportunidade"]) + bonus)
        l["Score de Oportunidade"] = sc; l["Score Comercial"] = sc; l["Prioridade"] = atk.faixa(sc)
        if bonus: l["Motivo da Prioridade"] = l["Motivo da Prioridade"].rstrip(".") + f"; fornecedor recorrente ({rec} contratos)."
        l["Responsável Comercial Sugerido"] = atribuir_responsavel(l, cfg)
        l["Motivo da Ação"] = f"{l['Urgência Comercial']} · {l['Janela Comercial']}: {l['Próxima Ação Recomendada']}"

def gerar_calibracao(path, L, agg, cfg, params, prev_date):
    rev = [l for l in L if l["Necessita Revisão?"] == "Sim" or l["Confiança da Classificação"] == "Baixa" or l["Fim da Vigência"] == ""]
    conc = sorted(agg["concorrentes"])
    m = ["# ATLAS B2G — Relatório de Calibração (Primeira Rodada de Produção)\n",
         f"**Rodada {params['data']} · fonte PNCP (dados públicos)**\n",
         "## 1. O que foi ajustado para produção",
         f"- **Janelas relativas** (capturam contratos vencendo agora): {params['janelas']}.",
         f"- **Concorrentes conhecidos** cadastrados: {len(cfg.get('concorrentes_conhecidos',[]))} (marcam Ameaça Alta e elevam o score). *A validar pela empresa.*",
         "- **Responsável por UF × categoria** ativo (cyber / software / serviços / infra / default).",
         f"- **Score mínimo** de corte: {cfg.get('score_minimo')}. **Bônus de recorrência** do fornecedor aplicado ao score.",
         f"- **Cobertura**: até {cfg.get('max_paginas_por_janela')} páginas por janela ({len(params['janelas'])} janelas).",
         "\n## 2. Parâmetros usados"]
    for k, v in params.items(): m.append(f"- {k}: {v}")
    m += ["\n## 3. Resultado",
          f"- Oportunidades: **{agg['n']}** · críticas: **{agg['criticas']}** · valor: **{utils.brl(agg['valor'])}** · DF {agg['uf'].get('DF',0)} / GO {agg['uf'].get('GO',0)}.",
          f"- Possíveis concorrentes detectados: {', '.join(conc[:12]) or '—'}.",
          "\n## 4. Limitações que permanecem",
          "- Janela por **data de publicação** (proxy de vigência) — cobertura amostral; aumente `max_paginas_por_janela`/janelas para ampliar.",
          "- 'Possível concorrente' é heurístico/por lista — **requer confirmação da empresa**.",
          "- Classificação por palavras-chave — objetos ambíguos sobram (ver Revisão Manual).",
          "- 'Servidor' tratado de forma conservadora (pode omitir compras de servidores sem qualificador técnico).",
          "\n## 5. O que precisa de validação humana",
          f"- **{len(rev)}** itens marcados *Necessita Revisão / baixa confiança / sem data*.",
          "- Confirmar os concorrentes reais e ajustar a lista em `config/atlas_config_producao.json`.",
          "- Confirmar renovação/aditivo nos contratos **vencidos** antes de abordar.",
          "\n## 6. O que a coordenação deve revisar",
          "- Nomes reais dos responsáveis no config e a distribuição sugerida (ver Mapa de Distribuição).",
          "- Top 10 da semana e alvos críticos de alto valor.",
          "- Itens da aba **Revisão Manual** do Excel antes de liberar para abordagem.",
          "\n---\n*Nenhuma oportunidade é garantida; a base depende de validação humana. Dados públicos do PNCP.*"]
    open(path, "w", encoding="utf-8").write("\n".join(m))

def gerar_mapa_distribuicao(md_path, csv_path, L, agg, cfg):
    porr = collections.defaultdict(list)
    for l in L: porr[l["Responsável Comercial Sugerido"]].append(l)
    linhas_csv = []
    for resp, ops in sorted(porr.items(), key=lambda x: -sum(o["Valor Total"] for o in x[1])):
        top = sorted(ops, key=lambda x: x["Score de Oportunidade"], reverse=True)[0]
        linhas_csv.append({"Responsável": resp, "Oportunidades": len(ops),
            "Críticas": sum(1 for o in ops if o["Urgência Comercial"] == "Crítica"),
            "Valor Total": round(sum(o["Valor Total"] for o in ops)),
            "UFs": ", ".join(sorted(set(o["UF"] for o in ops))), "Top Órgão": top["Nome Padronizado (Órgão)"]})
    utils.escrever_csv(csv_path, linhas_csv, ["Responsável", "Oportunidades", "Críticas", "Valor Total", "UFs", "Top Órgão"])
    m = [f"# ATLAS B2G — Mapa de Distribuição Comercial\n", f"**{agg['n']} oportunidades · {utils.brl(agg['valor'])} · críticas {agg['criticas']}**\n",
         "## Por responsável sugerido", "| Responsável | Oport. | Críticas | Valor | UFs |", "|---|---|---|---|---|"]
    for r in linhas_csv: m.append(f"| {r['Responsável']} | {r['Oportunidades']} | {r['Críticas']} | {utils.brl(r['Valor Total'])} | {r['UFs']} |")
    m.append("\n## Por UF")
    for uf in cfg["ufs"]: m.append(f"- **{uf}**: {agg['uf'].get(uf,0)} oport. · {utils.brl(agg['uf_valor'].get(uf,0))}")
    m.append("\n## Por categoria")
    for k, v in agg["cat"].most_common(12): m.append(f"- {k}: {v} · {utils.brl(agg['cat_valor'].get(k,0))}")
    m.append("\n## Valor por carteira")
    cart = collections.defaultdict(float)
    for l in L: cart[l["Carteira Sugerida"]] += l["Valor Total"]
    for k, v in sorted(cart.items(), key=lambda x: -x[1]): m.append(f"- {k}: {utils.brl(v)}")
    m.append("\n## Top 10 da semana (para distribuir)")
    for i, l in enumerate(L[:10], 1):
        m.append(f"{i}. **{l['Nome Padronizado (Órgão)']}** ({l['UF']}) · {l['Categoria Principal']} · {utils.brl(l['Valor Total'])} · {l['Prioridade']} → **{l['Responsável Comercial Sugerido']}** — {l['Próxima Ação Recomendada']}")
    open(md_path, "w", encoding="utf-8").write("\n".join(m))

def _notificar_ams(dburl, rodada_id):
    """1 notificação-resumo por AM/SE/Intern: novos contratos + a vencer (≤6 meses) nos órgãos dele.
    Sem spam (um resumo por coleta). Só Postgres; nunca quebra a rodada."""
    if not dburl or not rodada_id:
        return
    import psycopg2
    cn = psycopg2.connect(dburl); cn.autocommit = True; c = cn.cursor()
    try:
        c.execute("SELECT max(rodada_id) FROM oportunidades WHERE rodada_id < %s", (rodada_id,))
        prior = c.fetchone()[0]
        c.execute("""SELECT p.user_id, u.email, array_agg(po.orgao_id)
                     FROM perfis p JOIN auth.users u ON u.id = p.user_id
                     JOIN perfil_orgaos po ON po.user_id = p.user_id
                     WHERE p.ativo AND p.role IN ('Account Manager','Sales Engineer','Intern')
                     GROUP BY p.user_id, u.email""")
        ams = c.fetchall()
        feitas = 0
        for uid, email, orgaos in ams:
            c.execute("""SELECT count(*) FROM oportunidades op JOIN contratos ct ON ct.id = op.contrato_id
                         WHERE op.rodada_id = %s AND op.orgao_id = ANY(%s)
                         AND (%s IS NULL OR NOT EXISTS (
                             SELECT 1 FROM oportunidades o2 JOIN contratos c2 ON c2.id = o2.contrato_id
                             WHERE o2.rodada_id = %s AND c2.id_pncp = ct.id_pncp))""",
                      (rodada_id, orgaos, prior, prior))
            novos = c.fetchone()[0]
            c.execute("""SELECT count(*) FROM oportunidades op JOIN contratos ct ON ct.id = op.contrato_id
                         WHERE op.rodada_id = %s AND op.orgao_id = ANY(%s)
                         AND ct.dias_ate_vencimento BETWEEN 0 AND 180""",
                      (rodada_id, orgaos))
            vencendo = c.fetchone()[0]
            if novos == 0 and vencendo == 0:
                continue
            titulo = (f"{novos} novo(s) contrato(s) nos seus órgãos" if novos
                      else f"{vencendo} contrato(s) vencendo em ≤6 meses")
            msg = (f"Nos órgãos que você acompanha: {novos} novo(s) contrato(s) detectado(s) e "
                   f"{vencendo} vencendo em até 6 meses. Abra a Lista de Ataque para priorizar a prospecção.")
            nivel = "warning" if vencendo > 0 else "info"
            c.execute("""INSERT INTO notificacoes (tipo,titulo,mensagem,nivel,usuario_destino_id,usuario_destino_email,escopo,link_url,criada_por,rodada_id)
                         VALUES ('base_novos',%s,%s,%s,%s,%s,'user','/lista-ataque','coleta',%s)""",
                      (titulo, msg, nivel, uid, email, rodada_id))
            feitas += 1
        log(f"-> notificações por órgão: {feitas} pessoa(s) avisada(s)")
    finally:
        cn.close()

# ---------------- MAIN ----------------
def main():
    t0 = time.time(); args = parse_args()
    cfg = aplicar_config(utils.carregar_config(args.config), args)
    run_date = datetime.datetime.strptime(args.date, "%Y-%m-%d").date() if args.date else datetime.date.today()
    DATA = run_date.strftime("%Y-%m-%d")
    TAG = (args.tag.strip() if args.tag else "")
    folder = (TAG + "_" if TAG else "") + DATA
    round_dir = P("outputs", "rodadas", folder); os.makedirs(round_dir, exist_ok=True); os.makedirs(P("logs"), exist_ok=True)
    log(f"=== ATLAS B2G — rodada {folder}{' [PRODUÇÃO]' if TAG else ''} ===")
    # LOCK: impede duas coletas simultâneas (botão + automática + direta) que dariam throttle/conflito.
    # Robusto: verifica se o PID dono do lock ainda existe — se a coleta morreu, o lock não trava o sistema.
    import atexit
    lock_path = P("logs", ".scan.lock")
    def _pid_viva(pid):
        try:
            if os.name == "nt":
                import ctypes
                h = ctypes.windll.kernel32.OpenProcess(0x1000, False, int(pid))
                if h: ctypes.windll.kernel32.CloseHandle(h); return True
                return False
            os.kill(int(pid), 0); return True
        except Exception:
            return False
    try:
        if os.path.exists(lock_path):
            try: dono = int(open(lock_path, encoding="utf-8").read().split()[0])
            except Exception: dono = None
            recente = (time.time() - os.path.getmtime(lock_path)) < 13 * 3600
            if dono and recente and _pid_viva(dono):
                log(f"Já existe uma coleta em andamento (PID {dono}). Encerrando para não conflitar.")
                return
            log("Lock anterior obsoleto (processo não existe mais) — assumindo.")
        with open(lock_path, "w", encoding="utf-8") as _lf:
            _lf.write(f"{os.getpid()} {DATA}")
        atexit.register(lambda: os.path.exists(lock_path) and os.remove(lock_path))
    except Exception as _e:
        log(f"aviso: lock não criado ({_e}); seguindo mesmo assim.")
    # configura pipeline (cache em data/.cache_pncp)
    core.CONFIG["OUT_DIR"] = P("data"); core.CONFIG["USAR_CACHE"] = cfg["usar_cache"]
    core.CONFIG["RETRIES"] = cfg["http"]["retries"]; core.CONFIG["TIMEOUT"] = cfg["http"]["timeout"]; core.CONFIG["SLEEP"] = cfg["http"]["sleep"]
    core.CONFIG["CONCORRENTES_CONHECIDOS"] = cfg["concorrentes_conhecidos"]
    if cfg.get("limpar_cache"):
        cdir = P("data", ".cache_pncp"); shutil.rmtree(cdir, ignore_errors=True); log("cache limpo.")
    janelas = montar_janelas(cfg, run_date)
    atk.ATK.update({"UFS": cfg["ufs"], "JANELAS": janelas, "MAX_PAG_POR_JANELA": cfg["max_paginas_por_janela"],
        "VENCIDO_RECENTE_DIAS": cfg["vencido_recente_dias"], "FUTURO_MAX_DIAS": cfg["futuro_max_dias"],
        "MONITOR_MAX_DIAS": cfg["monitor_max_dias"], "VALOR_ALTO": cfg["valor_alto"], "OUT_DIR": round_dir,
        "CONCORRENTES": cfg["concorrentes_conhecidos"], "CARTEIRA_POR_UF": cfg["carteira_por_uf"], "RESP_POR_UF": cfg["responsavel_por_uf"]})
    log(f"UFs={cfg['ufs']} | janelas={janelas} | score_min={cfg['score_minimo']} | cache={cfg['usar_cache']}")
    # ETAPA 2-3: ingestão + lista
    try:
        regs = atk.coletar()
        L = atk.processar(regs)
    except Exception as e:
        log(f"FALHA na ingestão: {e}"); _salvar_log(round_dir, DATA); return
    enriquecer_producao(L, cfg)
    L = [l for l in L if l["Score de Oportunidade"] >= cfg["score_minimo"]]
    L.sort(key=lambda x: (x["Score de Oportunidade"], x["Valor Total"]), reverse=True)
    for n, l in enumerate(L, 1): l["ID Oportunidade"] = f"ATK-{n:05d}"
    if not L:
        log("Sem oportunidades após filtros. Encerrando."); _salvar_log(round_dir, DATA); return
    agg = atk.agregar(L)
    agg["imediato"] = sum(1 for l in L if l.get("Janela Comercial") == "Ataque imediato")
    # comparativo (mesmo rótulo)
    prev_date, removidas = comparar(L, P("outputs", "rodadas"), DATA, TAG or None)
    log(f"Oportunidades: {agg['n']} | críticas: {agg['criticas']} | valor: {utils.brl(agg['valor'])} | anterior: {prev_date or '—'}")
    # ETAPA 3: CSVs
    arquivos = []
    cols = list(L[0].keys())
    csv_lista = os.path.join(round_dir, f"atlas_lista_ataque_comercial_{DATA}.csv")
    utils.escrever_csv(csv_lista, L, cols); arquivos.append(csv_lista); log(f"-> {os.path.basename(csv_lista)}")
    csv_vend = os.path.join(round_dir, f"atlas_top_oportunidades_vendedores_{DATA}.csv")
    atk.vendedores_csv(csv_vend, L); arquivos.append(csv_vend)
    params = {"data": DATA, "ufs": cfg["ufs"], "janelas": janelas, "score_minimo": cfg["score_minimo"],
              "max_paginas_por_janela": cfg["max_paginas_por_janela"], "concorrentes": cfg["concorrentes_conhecidos"] or "(nenhum)",
              "fonte": "PNCP (dados públicos)", "versao_pipeline": cfg.get("_versao_pipeline", "ATLAS B2G v1.0"),
              "executado_em": f"{datetime.datetime.now():%d/%m/%Y %H:%M}"}
    # ETAPA 4: Excel
    if cfg["gerar_excel"]:
        xlsx = os.path.join(round_dir, f"atlas_lista_ataque_comercial_{DATA}.xlsx")
        try: gerar_excel(xlsx, L, agg, removidas, prev_date, params); arquivos.append(xlsx); log(f"-> {os.path.basename(xlsx)}")
        except Exception as e: log(f"Excel falhou: {e}")
    # ETAPA 6: Relatório
    rel_md = os.path.join(round_dir, f"atlas_relatorio_semanal_{DATA}.md")
    if cfg["gerar_relatorio"]:
        pdf, okpdf = gerar_relatorio(rel_md, L, agg, removidas, prev_date, params, time.time() - t0)
        arquivos.append(rel_md); log(f"-> {os.path.basename(rel_md)}")
        if okpdf and pdf: arquivos.append(pdf); log(f"-> {os.path.basename(pdf)}")
        else: log("PDF não gerado (reportlab ausente) — use o .md/.xlsx ou imprima o protótipo.")
    # ETAPA 5: Protótipo
    if cfg["gerar_prototipo"]:
        html = os.path.join(round_dir, f"ATLAS-B2G_Prototipo-Dados-Reais-PNCP_{DATA}.html")
        js = os.path.join(round_dir, f"atlas_dados_reais_prototipo_{DATA}.js")
        try: gerar_prototipo(L, html, js); arquivos += [html, js]; log(f"-> {os.path.basename(html)}")
        except Exception as e: log(f"Protótipo falhou: {e}")
    # EXTRAS de produção: relatório de calibração + mapa de distribuição
    if TAG:
        calib = os.path.join(round_dir, f"atlas_relatorio_calibracao_{DATA}.md")
        try: gerar_calibracao(calib, L, agg, cfg, params, prev_date); arquivos.append(calib); log(f"-> {os.path.basename(calib)}")
        except Exception as e: log(f"calibração falhou: {e}")
        mapa_md = os.path.join(round_dir, f"mapa_distribuicao_comercial_{DATA}.md")
        mapa_csv = os.path.join(round_dir, f"mapa_distribuicao_comercial_{DATA}.csv")
        try: gerar_mapa_distribuicao(mapa_md, mapa_csv, L, agg, cfg); arquivos += [mapa_md, mapa_csv]; log(f"-> {os.path.basename(mapa_md)}")
        except Exception as e: log(f"mapa falhou: {e}")
    # resumo executivo .txt
    resumo = os.path.join(round_dir, f"atlas_resumo_executivo_{DATA}.txt")
    _resumo_txt(resumo, L, agg, prev_date, removidas, DATA); arquivos.append(resumo)
    # cópia da configuração usada (para o pacote)
    try:
        cfg_copy = os.path.join(round_dir, os.path.basename(args.config)); shutil.copyfile(args.config, cfg_copy); arquivos.append(cfg_copy)
    except Exception: pass
    # log
    logf = _salvar_log(round_dir, DATA); arquivos.append(logf)
    # ETAPA 7: pacote ZIP
    if cfg["compactar_entrega"]:
        zipbase = f"atlas_pacote_{'producao' if TAG.upper()=='PRODUCAO' else 'envio_comercial'}_{DATA}.zip"
        zp = os.path.join(round_dir, zipbase)
        utils.zipar(zp, arquivos); log(f"-> {os.path.basename(zp)} ({len(arquivos)} arquivos)")
    # ETAPA 8: gravar no banco (opcional)
    if getattr(args, "write_db", False):
        try:
            import load_weekly_to_db as dbload
            res = dbload.carregar(rodada=round_dir, db_url=args.db_url, init=not getattr(args, "no_init", False), verbose=False, remove=not getattr(args, "no_remove", False))
            log(f"-> banco [{res['dialect']}]: rodada_id={res['rodada_id']} · oportunidades {res['oportunidades_novas']} novas / {res['oportunidades_total']} no total · histórico {res['historico_total']} · removidas {res['removidas']}")
            # CRÍTICO: a carga pode recriar as views e zerar o security_invoker, o que vazaria
            # a RLS por órgão (todos veriam tudo). Reforça após cada gravação (apenas Postgres).
            if res.get("dialect") == "postgres":
                _dburl = args.db_url or os.environ.get("DATABASE_URL")
                try:
                    import psycopg2 as _pg
                    _cn = _pg.connect(_dburl); _cn.autocommit = True; _cur = _cn.cursor()
                    for _v in ("vw_oportunidades_ativas", "vw_lista_ataque_atual", "vw_top_10_semana",
                               "vw_oportunidades_por_responsavel", "vw_concorrentes", "vw_qualidade_base",
                               "vw_dashboard_executivo", "vw_historico_rodadas"):
                        try: _cur.execute(f"ALTER VIEW {_v} SET (security_invoker = on)")
                        except Exception: pass
                    _cn.close()
                    log("-> security_invoker reforçado nas views (escopo por órgão preservado)")
                except Exception as _e:
                    log(f"AVISO: não reforçou security_invoker ({_e}) — rode rls_policies se necessário")
                try:
                    _notificar_ams(_dburl, res.get("rodada_id"))
                except Exception as _e:
                    log(f"AVISO: notificações por órgão não geradas ({_e})")
        except Exception as e:
            log(f"gravação no banco falhou (arquivos OK): {e}")
    log(f"=== concluído em {time.time()-t0:.1f}s · pasta: {round_dir} ===")
    _salvar_log(round_dir, DATA)
    print("\nRESUMO:", agg["n"], "oportunidades |", agg["criticas"], "críticas |", utils.brl(agg["valor"]), "| pasta:", round_dir)

def _resumo_txt(path, L, agg, prev_date, removidas, DATA):
    out = [f"ATLAS B2G — Resumo Executivo da Rodada {DATA}", "=" * 50,
           f"Oportunidades: {agg['n']} | Críticas: {agg['criticas']} | Valor total: {utils.brl(agg['valor'])}",
           f"Ataque imediato: {agg.get('imediato',0)} | Vencidos: {agg['vencidos']} | Vencendo <=30d: {agg['v30']}",
           f"DF: {agg['uf'].get('DF',0)} | GO: {agg['uf'].get('GO',0)}",
           f"Comparativo: {'rodada anterior '+prev_date if prev_date else 'primeira rodada'} | removidas: {len(removidas)}",
           "", "TOP 10 DA SEMANA:"]
    for i, l in enumerate(L[:10], 1):
        out.append(f"{i:2d}. [{l['Score de Oportunidade']}|{l['Urgência Comercial']}] {l['Nome Padronizado (Órgão)']} ({l['UF']}) · {l['Categoria Principal']} · {utils.brl(l['Valor Total'])} · {l['Status do Contrato']}")
        out.append(f"     Ação: {l['Próxima Ação Recomendada']}")
    out += ["", "Dados públicos do PNCP. Itens 'necessita revisão' dependem de validação humana. Nenhuma oportunidade é garantida."]
    open(path, "w", encoding="utf-8").write("\n".join(out))

def _salvar_log(round_dir, DATA):
    p = os.path.join(round_dir, f"atlas_log_execucao_{DATA}.txt")
    full = "PARÂMETROS:\n" + "\n".join(f"  {k}={v}" for k, v in atk.ATK.items()) + "\n\nLOG ORQUESTRADOR:\n" + "\n".join(RUN_LOG) + "\n\nLOG PIPELINE:\n" + "\n".join(atk._LOG)
    open(p, "w", encoding="utf-8").write(full)
    # cópia em logs/
    try: open(os.path.join(P("logs"), f"run_{DATA}.txt"), "w", encoding="utf-8").write(full)
    except Exception: pass
    return p

if __name__ == "__main__":
    main()
